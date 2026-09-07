import asyncio
import os
from datetime import date, datetime, timedelta
from io import BytesIO

os.environ["CMH_SMS_DATABASE_URL"] = "sqlite:///./test_serial_management.db"
os.environ["CMH_SMS_AUDIO_ENABLED"] = "false"

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.announcement import AnnouncementEngine
from app.auth import hash_password
from app.database import Base, engine
from app.main import app
from app.models import AppSetting, Doctor, QueueToken, RoleDefinition, SmsMessage, User, WaitingRoom
from app.notification import SmsOutboxWorker

client = TestClient(app)


def setup_function():
    client.cookies.clear()
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            WaitingRoom(
                id="waiting-room-1",
                code="WR-1",
                name="Medicine Waiting Room",
                floor="2nd Floor",
                display_label="Waiting Room 1",
            )
        )
        db.add(
            Doctor(
                id="dr-khan",
                name="Dr. Ayesha Khan",
                department="Medicine",
                designation="Consultant",
                room_number="205",
                waiting_room_id="waiting-room-1",
                token_prefix="AK",
            )
        )
        db.add(
            User(
                username="admin",
                full_name="Test Admin",
                password_hash=hash_password("AdminPass123!"),
                role="admin",
                is_active=True,
            )
        )
        db.commit()
    client.post("/api/v1/auth/login", json={"username": "admin", "password": "AdminPass123!"}).raise_for_status()


def token_payload(name="Rahim Uddin"):
    return {
        "patient_name": name,
        "patient_phone": "01700000000",
        "service_category": "army",
        "rank": "Captain",
        "service_number": "BA-TEST-001",
        "doctor_id": "dr-khan",
        "doctor_name": "Dr. Ayesha Khan",
        "department": "Medicine",
        "room_number": "205",
        "waiting_room": "WR-1",
        "source": "walk_in",
        "priority": "normal",
    }


def test_short_invalid_login_returns_authentication_error_not_validation_error():
    response = client.post("/api/v1/auth/login", json={"username": "x", "password": "bad"})

    assert response.status_code == 401
    assert response.json() == {"detail": "Invalid username or password"}


def test_admin_can_update_doctor_room_number():
    response = client.patch("/api/v1/admin/doctors/dr-khan", json={"room_number": "207"})

    assert response.status_code == 200
    assert response.json()["room_number"] == "207"
    assert client.get("/api/v1/doctors").json()[0]["room"] == "207"


def test_waiting_patient_room_can_change_but_called_patient_room_is_locked():
    token = client.post("/api/v1/tokens", json=token_payload()).json()

    updated = client.patch(f"/api/v1/tokens/{token['id']}/room", json={"room_number": "207"})

    assert updated.status_code == 200
    assert updated.json()["room_number"] == "207"
    called = client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").json()
    assert called["room_number"] == "207"
    locked = client.patch(f"/api/v1/tokens/{token['id']}/room", json={"room_number": "208"})
    assert locked.status_code == 409
    assert "after the patient has been called" in locked.json()["detail"]


def test_reception_to_display_flow():
    first = client.post("/api/v1/tokens", json=token_payload()).json()
    second = client.post("/api/v1/tokens", json=token_payload("Karim Ahmed")).json()
    assert first["token_number"] == "AK-001"
    assert second["token_number"] == "AK-002"

    called = client.post("/api/v1/doctors/dr-khan/call-next").json()
    assert called["id"] == first["id"]
    assert called["status"] == "called"

    display = client.get("/api/v1/displays/WR-1").json()
    assert display["current"]["patient_name"] == "Rahim Uddin"
    assert len(display["active_calls"]) == 1
    assert display["next_tokens"][0]["patient_name"] == "Karim A."
    assert "room 205" in display["announcement"]
    assert "Doctor Dr." not in display["announcement"]


def test_reception_report_filters_and_exports_complete_excel_table():
    client.post("/api/v1/tokens", json=token_payload()).raise_for_status()
    priority = token_payload("Priority Patient")
    priority["priority"] = "priority"
    client.post("/api/v1/tokens", json=priority).raise_for_status()
    params = {
        "date_from": date.today().isoformat(),
        "date_to": date.today().isoformat(),
        "doctor_id": "dr-khan",
        "waiting_room": "WR-1",
        "priority": "priority",
        "service_category": "army",
        "status": "waiting",
    }

    report = client.get("/api/v1/reports/reception", params=params)
    assert report.status_code == 200
    assert [row["patient_name"] for row in report.json()] == ["Priority Patient"]
    assert report.json()[0]["token_date"] == date.today().isoformat()

    exported = client.get("/api/v1/reports/reception.xlsx", params=params)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook["Reception Desk"]
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.freeze_panes == "F2"
    assert sheet["E2"].value == "Priority Patient"
    assert sheet.max_column == 37
    assert sheet.max_row == 2

    pdf = client.get("/api/v1/reports/reception.pdf", params=params)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content.startswith(b"%PDF")


def test_seeded_bilingual_setting_is_understood_by_server_audio():
    engine = AnnouncementEngine()
    engine.enabled = True
    captured = []
    engine._queue = type("CaptureQueue", (), {"put_nowait": lambda self, item: captured.append(item)})()
    token = type("Token", (), {"token_number": "AK-001", "patient_name": "Rahim Uddin", "room_number": "205"})()

    assert engine.enqueue(token, {"language_order": ["bn", "en"], "repeat_count": 2, "rate": 0.88})
    assert captured[0].language == "both"
    assert captured[0].repeat_count == 2


def test_doctor_call_and_recall_are_sent_to_server_speaker(monkeypatch):
    queued = []
    monkeypatch.setattr(
        "app.main.announcement_engine.enqueue", lambda token, settings: queued.append(token.token_number) or True
    )
    token = client.post("/api/v1/tokens", json=token_payload()).json()

    client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").raise_for_status()
    client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action",
        json={"action": "recall", "actor": "doctor.test"},
    ).raise_for_status()

    assert queued == ["AK-001", "AK-001"]


def test_macos_say_generates_aiff(monkeypatch, tmp_path):
    engine = AnnouncementEngine()
    engine.cache_dir = tmp_path
    commands = []
    monkeypatch.setattr(engine, "_synthesizer", lambda: "/usr/bin/say")
    monkeypatch.setattr(engine, "_macos_voice", lambda language: "Samantha" if language == "en" else None)
    monkeypatch.setattr("app.announcement.subprocess.run", lambda command, **kwargs: commands.append(command))

    output = engine._synthesise("Token AK-001", "en", 150)

    assert output.suffix == ".aiff"
    assert commands == [["/usr/bin/say", "-v", "Samantha", "-r", "150", "-o", str(output), "Token AK-001"]]


def test_macos_bangla_uses_installed_bengali_voice(monkeypatch, tmp_path):
    engine = AnnouncementEngine()
    engine.cache_dir = tmp_path
    commands = []
    monkeypatch.setattr(engine, "_synthesizer", lambda: "/usr/bin/say")
    monkeypatch.setattr(engine, "_macos_voice", lambda language: "Piya" if language == "bn" else "Samantha")
    monkeypatch.setattr("app.announcement.subprocess.run", lambda command, **kwargs: commands.append(command))

    output = engine._synthesise("টোকেন নম্বর AK-001", "bn", 150)

    assert commands == [["/usr/bin/say", "-v", "Piya", "-r", "150", "-o", str(output), "টোকেন নম্বর AK-001"]]


def test_espeak_bangla_support_is_reported(monkeypatch):
    completed = type("Completed", (), {"stdout": " 5  bn              M  bengali             bn"})()
    monkeypatch.setattr("app.announcement.subprocess.run", lambda command, **kwargs: completed)

    assert AnnouncementEngine._bangla_supported("/usr/bin/espeak-ng")


def test_bangla_announcement_is_one_continuous_voice_clip(monkeypatch, tmp_path):
    engine = AnnouncementEngine()
    engine.prompt_dir = tmp_path
    intro = tmp_path / "bangla-intro.wav"
    outro = tmp_path / "bangla-outro.wav"
    intro.touch()
    outro.touch()
    captured = []
    monkeypatch.setattr(
        engine,
        "_synthesise",
        lambda text, voice, rate: captured.append((voice, text)) or tmp_path / f"dynamic-{len(captured)}.wav",
    )
    item = type(
        "AnnouncementItem",
        (),
        {
            "language": "both",
            "token_number": "AK-001",
            "patient_name": "Rahim Uddin",
            "room_number": "205",
            "rate": 150,
        },
    )()

    paths = engine._prepare(item)

    assert len(paths) == 2
    assert captured == [
        ("bn_female", "রোগী Rahim Uddin, অনুগ্রহ করে কক্ষ নম্বর দুই শূন্য পাঁচ-এ যান।"),
        ("en_male", "Rahim Uddin, please proceed to your radiographer, room 205."),
    ]


def test_bangla_generation_does_not_depend_on_prerecorded_prompts(monkeypatch, tmp_path):
    engine = AnnouncementEngine()
    engine.prompt_dir = tmp_path
    captured = []
    monkeypatch.setattr(
        engine,
        "_synthesise",
        lambda text, voice, rate: captured.append((voice, text)) or tmp_path / "english.wav",
    )
    item = type(
        "AnnouncementItem",
        (),
        {
            "language": "both",
            "token_number": "AK-001",
            "patient_name": "Rahim Uddin",
            "doctor_name": "Dr. Ayesha Khan",
            "room_number": "205",
            "rate": 150,
        },
    )()

    engine._prepare(item)

    assert captured == [
        ("bn_female", "রোগী Rahim Uddin, অনুগ্রহ করে কক্ষ নম্বর দুই শূন্য পাঁচ-এ যান।"),
        ("en_male", "Rahim Uddin, please proceed to Dr. Ayesha Khan, room 205.")
    ]


def test_multiple_doctors_share_one_fifo_audio_queue():
    engine = AnnouncementEngine()
    engine.enabled = True
    first = type("Token", (), {"token_number": "AK-001", "patient_name": "Rahim", "room_number": "205"})()
    second = type("Token", (), {"token_number": "FR-001", "patient_name": "Karim", "room_number": "312"})()
    third = type("Token", (), {"token_number": "NS-001", "patient_name": "Nusrat", "room_number": "118"})()

    engine.enqueue(first, {"language": "en"})
    engine.enqueue(second, {"language": "en"})
    engine.enqueue(third, {"language": "en"})

    assert engine.status()["queued"] == 3
    assert [engine._queue.get_nowait().token_number for _ in range(3)] == ["AK-001", "FR-001", "NS-001"]


def test_audio_worker_finishes_each_call_before_starting_next(monkeypatch):
    async def scenario():
        engine = AnnouncementEngine()
        engine.enabled = True
        order = []
        monkeypatch.setattr(engine, "_prepare", lambda item: [item.token_number])
        monkeypatch.setattr(engine, "_play", lambda path: order.append(path))
        await engine.start()
        try:
            for token_number in ("AK-001", "FR-001", "NS-001"):
                token = type(
                    "Token", (), {"token_number": token_number, "patient_name": "Patient", "room_number": "1"}
                )()
                engine.enqueue(token, {"language": "en"})
            await asyncio.wait_for(engine._queue.join(), timeout=2)
        finally:
            await engine.stop()
        return order, engine.completed_count

    order, completed = asyncio.run(scenario())
    assert order == ["AK-001", "FR-001", "NS-001"]
    assert completed == 3


def test_unauthenticated_api_is_rejected():
    client.cookies.clear()
    session = client.get("/api/v1/auth/session")
    assert session.status_code == 200
    assert session.json() is None
    assert client.get("/api/v1/auth/me").status_code == 401
    assert client.get("/api/v1/tokens").status_code == 401
    assert client.get("/api/v1/settings").status_code == 401
    protected_gets = [
        "/api/v1/permissions",
        "/api/v1/users",
        "/api/v1/roles",
        "/api/v1/doctors",
        "/api/v1/waiting-rooms",
        "/api/v1/lookups",
        "/api/v1/holidays",
        "/api/v1/patients",
        "/api/v1/schedule-slots",
        "/api/v1/appointments",
        "/api/v1/dashboard",
        "/api/v1/audit",
        "/api/v1/reports/queue-summary",
        "/api/v1/reports/reception",
        "/api/v1/reports/reception.xlsx",
        "/api/v1/reports/reception.pdf",
        "/api/v1/devices",
        "/api/v1/notifications/sms",
        "/api/v1/displays/WR-1",
    ]
    for path in protected_gets:
        assert client.get(path).status_code == 401, path


def test_rbac_and_doctor_assignment_are_enforced():
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            WaitingRoom(
                id="waiting-room-2",
                code="WR-2",
                name="Surgery Waiting Room",
                floor="3rd Floor",
                display_label="Waiting Room 2",
            )
        )
        db.add(
            Doctor(
                id="dr-other",
                name="Dr. Other",
                department="Surgery",
                designation="Consultant",
                room_number="301",
                waiting_room_id="waiting-room-2",
                token_prefix="OT",
            )
        )
        db.add(
            User(
                username="reception",
                full_name="Reception",
                password_hash=hash_password("Reception123!"),
                role="reception",
                is_active=True,
            )
        )
        db.add(
            User(
                username="radiographer",
                full_name="Radiographer",
                password_hash=hash_password("Radiographer123!"),
                role="radiographer",
                doctor_id="dr-khan",
                is_active=True,
            )
        )
        db.commit()

    client.cookies.clear()
    client.post("/api/v1/auth/login", json={"username": "reception", "password": "Reception123!"}).raise_for_status()
    assigned_token = client.post("/api/v1/tokens", json=token_payload()).json()
    assigned_waiting = client.post("/api/v1/tokens", json=token_payload("Second Assigned Patient")).json()
    other_payload = token_payload("Other Queue Patient") | {
        "doctor_id": "dr-other",
        "doctor_name": "Dr. Other",
        "department": "Surgery",
        "room_number": "301",
        "waiting_room": "WR-2",
    }
    other_token = client.post("/api/v1/tokens", json=other_payload).json()
    assert client.post("/api/v1/doctors/dr-khan/call-next").status_code == 403
    assert client.get("/api/v1/settings").status_code == 403

    client.cookies.clear()
    client.post(
        "/api/v1/auth/login", json={"username": "radiographer", "password": "Radiographer123!"}
    ).raise_for_status()
    assert client.post("/api/v1/doctors/dr-other/call-next").status_code == 403
    assert client.post("/api/v1/doctors/dr-khan/call-next").status_code == 200
    room_update = client.patch(
        f"/api/v1/tokens/{assigned_waiting['id']}/room", json={"room_number": "209"}
    )
    assert room_update.status_code == 200
    assert room_update.json()["room_number"] == "209"
    assert client.patch(
        f"/api/v1/tokens/{assigned_token['id']}/room", json={"room_number": "210"}
    ).status_code == 409
    assert client.patch(
        f"/api/v1/tokens/{other_token['id']}/room", json={"room_number": "302"}
    ).status_code == 403
    scoped_dashboard = client.get("/api/v1/dashboard")
    assert scoped_dashboard.status_code == 200
    assert scoped_dashboard.json()["scope"] == "assigned"
    assert scoped_dashboard.json()["total"] == 3
    assert [row["doctor_id"] for row in scoped_dashboard.json()["radiographers"]] == ["dr-khan"]
    assert [doctor["id"] for doctor in client.get("/api/v1/doctors").json()] == ["dr-khan"]
    assert [room["code"] for room in client.get("/api/v1/waiting-rooms").json()] == ["WR-1"]
    assert {token["doctor_id"] for token in client.get("/api/v1/tokens").json()} == {"dr-khan", "dr-other"}
    assert client.get("/api/v1/displays/WR-1").status_code == 403
    assert client.post(
        f"/api/v1/tokens/{other_token['id']}/transfer",
        json={"doctor_id": "dr-khan", "reason": "Should not be allowed"},
    ).status_code == 403


def test_called_patient_uses_full_name_while_upcoming_names_remain_private():
    called = client.post("/api/v1/tokens", json=token_payload("MD. Atiqur Rahman")).json()
    client.post("/api/v1/tokens", json=token_payload("Karim Ahmed")).raise_for_status()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{called['id']}/call").raise_for_status()

    display = client.get("/api/v1/displays/WR-1").json()
    assert display["current"]["patient_name"] == "MD. Atiqur Rahman"
    assert "MD. Atiqur Rahman" in display["announcement"]
    assert display["next_tokens"][0]["patient_name"] == "Karim A."


def test_call_and_recall_broadcast_to_all_waiting_rooms():
    created = client.post("/api/v1/tokens", json=token_payload()).json()
    with client.websocket_connect("/api/v1/realtime/waiting-rooms/WR-1") as room_one:
        with client.websocket_connect("/api/v1/realtime/waiting-rooms/WR-2") as room_two:
            room_one.receive_json()
            room_two.receive_json()
            client.post(f"/api/v1/doctors/dr-khan/tokens/{created['id']}/call").raise_for_status()
            for socket in (room_one, room_two):
                called = socket.receive_json()
                assert called["type"] == "patient.called"
                assert called["display"]["current"]["id"] == created["id"]

            client.post(
                f"/api/v1/doctors/dr-khan/tokens/{created['id']}/action",
                json={"action": "recall", "actor": "doctor.test"},
            ).raise_for_status()
            for socket in (room_one, room_two):
                recalled = socket.receive_json()
                assert recalled["type"] == "patient.called"
                assert recalled["reason"] == "queue.recall"
                assert recalled["display"]["current"]["id"] == created["id"]


def test_global_call_remains_visible_when_another_room_refreshes():
    created = client.post("/api/v1/tokens", json=token_payload()).json()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{created['id']}/call").raise_for_status()

    other_room = client.get("/api/v1/displays/WR-3").json()
    assert other_room["current"]["id"] == created["id"]
    assert other_room["next_tokens"] == []


def test_waiting_room_can_hold_calls_for_multiple_doctors():
    from app.database import SessionLocal
    from app.models import Doctor

    with SessionLocal() as db:
        db.add(
            Doctor(
                id="dr-rahman",
                name="Dr. Farhan Rahman",
                department="Cardiology",
                designation="Consultant",
                room_number="312",
                waiting_room_id="waiting-room-1",
                token_prefix="FR",
            )
        )
        db.commit()
    second_payload = token_payload("Nusrat Jahan") | {
        "doctor_id": "dr-rahman",
        "doctor_name": "Dr. Farhan Rahman",
        "department": "Cardiology",
        "room_number": "312",
    }
    first = client.post("/api/v1/tokens", json=token_payload()).json()
    second = client.post("/api/v1/tokens", json=second_payload).json()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{first['id']}/call").raise_for_status()
    client.post(f"/api/v1/doctors/dr-rahman/tokens/{second['id']}/call").raise_for_status()
    display = client.get("/api/v1/displays/WR-1").json()
    assert {call["doctor_id"] for call in display["active_calls"]} == {"dr-khan", "dr-rahman"}


def test_doctor_state_transitions_are_validated_and_audited():
    token = client.post("/api/v1/tokens", json=token_payload()).json()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").raise_for_status()
    started = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "start", "actor": "doctor.test"}
    ).json()
    assert started["status"] == "in_progress"
    completed = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "complete", "actor": "doctor.test"}
    ).json()
    assert completed["status"] == "completed"
    invalid = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "recall", "actor": "doctor.test"}
    )
    assert invalid.status_code == 409
    actions = [event["action"] for event in client.get("/api/v1/audit").json()]
    assert "queue.start" in actions
    assert "queue.complete" in actions


def test_same_patient_can_be_announced_again_with_recall_limit():
    from app.database import SessionLocal
    from app.models import AppSetting

    with SessionLocal() as db:
        db.add(AppSetting(key="queue", value={"recall_limit": 2}, description="Queue rules"))
        db.commit()
    token = client.post("/api/v1/tokens", json=token_payload()).json()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").raise_for_status()
    first = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "recall", "actor": "doctor.test"}
    ).json()
    second = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "recall", "actor": "doctor.test"}
    ).json()
    assert first["recall_count"] == 1
    assert second["recall_count"] == 2
    blocked = client.post(
        f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={"action": "recall", "actor": "doctor.test"}
    )
    assert blocked.status_code == 409


def test_paused_queue_blocks_call_next_until_resumed():
    client.post("/api/v1/tokens", json=token_payload()).raise_for_status()
    client.put(
        "/api/v1/doctors/dr-khan/queue-control",
        json={
            "is_paused": True,
            "reason": "Doctor temporarily unavailable",
        },
    ).raise_for_status()
    blocked = client.post("/api/v1/doctors/dr-khan/call-next")
    assert blocked.status_code == 409
    client.put(
        "/api/v1/doctors/dr-khan/queue-control",
        json={
            "is_paused": False,
            "reason": "Doctor returned",
        },
    ).raise_for_status()
    assert client.post("/api/v1/doctors/dr-khan/call-next").status_code == 200


def test_reception_can_transfer_waiting_token_and_server_owns_doctor_metadata():
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            Doctor(
                id="dr-other",
                name="Dr. Real Name",
                department="Surgery",
                designation="Consultant",
                room_number="301",
                waiting_room_id="waiting-room-1",
                token_prefix="RN",
            )
        )
        db.commit()
    payload = token_payload() | {"doctor_name": "Tampered", "department": "Wrong", "room_number": "999"}
    token = client.post("/api/v1/tokens", json=payload).json()
    assert token["doctor_name"] == "Dr. Ayesha Khan"
    moved = client.post(
        f"/api/v1/tokens/{token['id']}/transfer",
        json={
            "doctor_id": "dr-other",
            "reason": "Requested specialty",
        },
    ).json()
    assert moved["doctor_id"] == "dr-other"
    assert moved["doctor_name"] == "Dr. Real Name"
    assert moved["token_number"] == "RN-001"


def test_device_credentials_protect_heartbeat_and_acknowledgement():
    created = client.post(
        "/api/v1/devices",
        json={
            "name": "Waiting Room One Display",
            "device_type": "display",
            "waiting_room": "WR-1",
        },
    ).json()
    device, key = created["device"], created["client_key"]
    assert client.post(f"/api/v1/devices/{device['id']}/heartbeat").status_code == 401
    heartbeat = client.post(
        f"/api/v1/devices/{device['id']}/heartbeat",
        headers={"X-Device-Key": key},
    )
    assert heartbeat.status_code == 200
    acknowledged = client.post(
        f"/api/v1/devices/{device['id']}/acknowledge",
        params={"event_id": "event-123"},
        headers={"X-Device-Key": key},
    ).json()
    assert acknowledged["last_acknowledged_event_id"] == "event-123"


def test_admin_configures_dropdown_values_and_inactive_values_are_rejected():
    created = client.post(
        "/api/v1/lookups",
        json={
            "category": "service_category",
            "value": "foreign_military",
            "label": "Foreign military",
            "sort_order": 20,
            "metadata_json": {},
        },
    ).json()
    visible = client.get("/api/v1/lookups", params={"category": "service_category"}).json()
    assert [(item["value"], item["label"]) for item in visible] == [("foreign_military", "Foreign military")]

    token = client.post("/api/v1/tokens", json=token_payload() | {"service_category": "foreign_military"})
    assert token.status_code == 201
    client.patch(f"/api/v1/lookups/{created['id']}", json={"is_active": False}).raise_for_status()
    blocked = client.post(
        "/api/v1/tokens", json=token_payload("Another Patient") | {"service_category": "foreign_military"}
    )
    assert blocked.status_code == 422
    assert client.get("/api/v1/lookups", params={"category": "service_category"}).json() == []


def test_patient_designation_comes_from_configured_master_data():
    client.post(
        "/api/v1/lookups",
        json={
            "category": "designation",
            "value": "capt",
            "label": "Capt.",
            "sort_order": 1,
            "metadata_json": {},
        },
    ).raise_for_status()
    token = client.post(
        "/api/v1/tokens",
        json=token_payload() | {"patient_title": "capt"},
    ).json()
    assert token["patient_title"] == "capt"
    invalid = client.post(
        "/api/v1/tokens",
        json=token_payload("Invalid Title") | {"patient_title": "invented"},
    )
    assert invalid.status_code == 422


def test_mobile_is_optional_and_sms_is_queued_when_present():
    without_mobile = client.post("/api/v1/tokens", json=token_payload() | {"patient_phone": ""})
    assert without_mobile.status_code == 201
    from sqlalchemy import select

    from app.database import SessionLocal
    from app.models import SmsMessage

    with SessionLocal() as db:
        assert db.scalar(select(SmsMessage).where(SmsMessage.mobile == "")) is None

    client.post("/api/v1/tokens", json=token_payload("SMS Patient")).raise_for_status()
    with SessionLocal() as db:
        message = db.scalar(select(SmsMessage).where(SmsMessage.mobile == "01700000000"))
        assert message.message_type == "token_confirmation"
        assert "room 205" in message.body
        assert message.status == "queued"


def test_readiness_and_security_headers_are_exposed():
    response = client.get("/api/v1/ready")
    assert response.status_code == 200
    assert response.json()["database"] == "ok"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    assert response.headers["x-request-id"]
    frontend = client.get("/")
    assert frontend.status_code == 200
    assert frontend.headers["cache-control"] == "no-store"


def test_unsafe_production_configuration_is_rejected(monkeypatch):
    import app.main as main

    monkeypatch.setattr(main, "ENVIRONMENT", "production")
    monkeypatch.setattr(main, "DATABASE_URL", "sqlite:///unsafe.db")
    monkeypatch.setattr(main, "COOKIE_SECURE", False)
    try:
        main.validate_production_config()
        assert False, "unsafe production configuration should fail"
    except RuntimeError as error:
        assert "PostgreSQL" in str(error)
        assert "COOKIE_SECURE" in str(error)


def test_builtin_administrator_role_cannot_be_downgraded():
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            RoleDefinition(
                name="admin",
                display_name="Administrator",
                access_profile="admin",
                description="Full administration",
                permissions=["*"],
                is_system=True,
            )
        )
        db.commit()

    response = client.patch("/api/v1/roles/admin", json={"access_profile": "reception"})

    assert response.status_code == 409
    with SessionLocal() as db:
        assert db.get(RoleDefinition, "admin").access_profile == "admin"


def test_role_assignment_permission_does_not_grant_account_update_permission():
    from app.database import SessionLocal

    client.cookies.clear()
    with SessionLocal() as db:
        db.add_all(
            [
                RoleDefinition(
                    name="role_manager",
                    display_name="Role manager",
                    access_profile="reception",
                    description="Can assign roles only",
                    permissions=["users.roles.assign"],
                    is_system=False,
                ),
                RoleDefinition(
                    name="reception",
                    display_name="Reception",
                    access_profile="reception",
                    description="Reception",
                    permissions=[],
                    is_system=True,
                ),
                User(
                    username="manager",
                    full_name="Role Manager",
                    password_hash=hash_password("ManagerPass123"),
                    role="role_manager",
                    is_active=True,
                ),
                User(
                    username="target",
                    full_name="Target User",
                    password_hash=hash_password("TargetPass123"),
                    role="reception",
                    is_active=True,
                ),
            ]
        )
        db.commit()
        target_id = db.scalar(select(User.id).where(User.username == "target"))

    client.post("/api/v1/auth/login", json={"username": "manager", "password": "ManagerPass123"}).raise_for_status()
    response = client.patch(
        f"/api/v1/users/{target_id}",
        json={"role": "reception", "is_active": False},
    )

    assert response.status_code == 403
    with SessionLocal() as db:
        assert db.get(User, target_id).is_active is True


def test_invalid_settings_are_rejected_without_corrupting_runtime_configuration():
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            AppSetting(
                key="display",
                value={"privacy_mode": "initials", "next_token_count": 5, "ticker_message": "Ready"},
                description="Display configuration",
            )
        )
        db.commit()

    response = client.put("/api/v1/settings/display", json={"value": {"next_token_count": "many"}})
    assert response.status_code == 422
    assert client.put("/api/v1/settings/unknown", json={"value": {}}).status_code == 404
    with SessionLocal() as db:
        assert db.get(AppSetting, "display").value["next_token_count"] == 5


def test_dashboard_counts_only_todays_tokens():
    from app.database import SessionLocal

    old = client.post("/api/v1/tokens", json=token_payload("Yesterday Patient")).json()
    current = client.post("/api/v1/tokens", json=token_payload("Today Patient")).json()
    with SessionLocal() as db:
        db.get(QueueToken, old["id"]).token_date = date.today() - timedelta(days=1)
        db.commit()

    client.post(f"/api/v1/doctors/dr-khan/tokens/{current['id']}/call").raise_for_status()
    dashboard = client.get("/api/v1/dashboard").json()

    assert dashboard["called"] == 1
    assert dashboard["waiting"] == 0


def test_transfer_refreshes_both_source_and_destination_waiting_rooms(monkeypatch):
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.add(
            WaitingRoom(
                id="waiting-room-2",
                code="WR-2",
                name="Surgery Waiting Room",
                floor="3rd Floor",
                display_label="Waiting Room 2",
            )
        )
        db.add(
            Doctor(
                id="dr-other",
                name="Dr. Other",
                department="Surgery",
                designation="Consultant",
                room_number="301",
                waiting_room_id="waiting-room-2",
                token_prefix="OT",
            )
        )
        db.commit()
    token = client.post("/api/v1/tokens", json=token_payload()).json()
    broadcasts = []

    async def capture(room, event_type, action, **payload):
        broadcasts.append((room, event_type, action, payload))

    monkeypatch.setattr("app.main.waiting_room_hub.broadcast", capture)
    client.post(
        f"/api/v1/tokens/{token['id']}/transfer",
        json={"doctor_id": "dr-other", "reason": "Requested specialty"},
    ).raise_for_status()

    assert [(room, action) for room, _, action, _ in broadcasts] == [
        ("WR-1", "queue.transferred_out"),
        ("WR-2", "queue.transferred"),
    ]


def test_patient_appointment_and_check_in_workflow_is_audited_and_idempotent():
    starts_at = (datetime.utcnow() + timedelta(days=1)).replace(hour=9, minute=0, second=0, microsecond=0)
    patient = client.post(
        "/api/v1/patients",
        json={"name": "Appointment Patient", "mobile": "01711111111"},
    )
    assert patient.status_code == 201
    patient_data = patient.json()
    slot = client.post(
        "/api/v1/schedule-slots",
        json={
            "doctor_id": "dr-khan",
            "starts_at": starts_at.isoformat(),
            "ends_at": (starts_at + timedelta(hours=1)).isoformat(),
            "capacity": 2,
        },
    )
    assert slot.status_code == 201
    appointment = client.post(
        "/api/v1/appointments",
        json={"patient_id": patient_data["id"], "doctor_id": "dr-khan", "slot_id": slot.json()["id"]},
    )
    assert appointment.status_code == 201
    appointment_data = appointment.json()
    confirmed = client.patch(
        f"/api/v1/appointments/{appointment_data['id']}",
        json={"action": "confirm", "reason": "Patient confirmed attendance"},
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "confirmed"

    first = client.post(f"/api/v1/appointments/{appointment_data['id']}/check-in")
    second = client.post(f"/api/v1/appointments/{appointment_data['id']}/check-in")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["id"] == second.json()["id"]
    assert first.json()["source"] == "appointment"

    events = client.get("/api/v1/audit?limit=20").json()
    actions = {event["action"] for event in events}
    assert {"patient.created", "schedule.slot_created", "appointment.created", "appointment.checked_in"} <= actions
    patient_event = next(event for event in events if event["action"] == "patient.created")
    assert patient_event["detail"]["patient_id"] == patient_data["id"]


def test_appointment_holiday_capacity_reschedule_and_cancel_rules():
    base = (datetime.utcnow() + timedelta(days=10)).replace(hour=8, minute=0, second=0, microsecond=0)

    def create_slot(starts_at: datetime, capacity: int = 1):
        return client.post(
            "/api/v1/schedule-slots",
            json={
                "doctor_id": "dr-khan",
                "starts_at": starts_at.isoformat(),
                "ends_at": (starts_at + timedelta(hours=1)).isoformat(),
                "capacity": capacity,
            },
        )

    first_patient = client.post(
        "/api/v1/patients", json={"name": "First Scheduled Patient", "mobile": "01722222222"}
    ).json()
    second_patient = client.post(
        "/api/v1/patients", json={"name": "Second Scheduled Patient", "mobile": "01733333333"}
    ).json()
    holiday_slot = create_slot(base).json()
    client.post(
        "/api/v1/holidays", json={"holiday_date": base.date().isoformat(), "name": "Hospital closure"}
    ).raise_for_status()
    holiday_booking = client.post(
        "/api/v1/appointments",
        json={"patient_id": first_patient["id"], "doctor_id": "dr-khan", "slot_id": holiday_slot["id"]},
    )
    assert holiday_booking.status_code == 409

    active_start = base + timedelta(days=1)
    active_slot = create_slot(active_start).json()
    first_booking = client.post(
        "/api/v1/appointments",
        json={"patient_id": first_patient["id"], "doctor_id": "dr-khan", "slot_id": active_slot["id"]},
    )
    first_booking.raise_for_status()
    full_booking = client.post(
        "/api/v1/appointments",
        json={"patient_id": second_patient["id"], "doctor_id": "dr-khan", "slot_id": active_slot["id"]},
    )
    assert full_booking.status_code == 409

    replacement = create_slot(active_start + timedelta(hours=2)).json()
    rescheduled = client.patch(
        f"/api/v1/appointments/{first_booking.json()['id']}",
        json={"action": "reschedule", "slot_id": replacement["id"], "reason": "Patient requested later time"},
    )
    assert rescheduled.status_code == 200
    assert rescheduled.json()["slot_id"] == replacement["id"]
    cancelled = client.patch(
        f"/api/v1/appointments/{first_booking.json()['id']}",
        json={"action": "cancel", "reason": "Patient no longer requires visit"},
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    assert client.post(f"/api/v1/appointments/{first_booking.json()['id']}/check-in").status_code == 409


def test_sms_manual_retry_resets_exhausted_attempts_and_worker_delivers(monkeypatch):
    from app.database import SessionLocal

    with SessionLocal() as db:
        message = SmsMessage(
            mobile="01744444444",
            message_type="token_confirmation",
            body="Test delivery",
            status="failed",
            attempts=5,
            error="Gateway unavailable",
            next_attempt_at=datetime.utcnow() + timedelta(days=1),
        )
        db.add(message)
        db.commit()
        message_id = message.id

    retried = client.post(f"/api/v1/notifications/sms/{message_id}/retry")
    assert retried.status_code == 200
    with SessionLocal() as db:
        reset = db.get(SmsMessage, message_id)
        assert reset.status == "queued"
        assert reset.attempts == 0
        assert reset.next_attempt_at is None

    class GatewayResponse:
        headers = {"content-type": "application/json"}

        @staticmethod
        def raise_for_status():
            return None

        @staticmethod
        def json():
            return {"message_id": "provider-123"}

    monkeypatch.setenv("CMH_SMS_SMS_GATEWAY_URL", "https://sms.invalid/send")
    monkeypatch.setattr("app.notification.httpx.post", lambda *args, **kwargs: GatewayResponse())
    SmsOutboxWorker._dispatch_batch()

    with SessionLocal() as db:
        sent = db.get(SmsMessage, message_id)
        assert sent.status == "sent"
        assert sent.attempts == 1
        assert sent.provider_reference == "provider-123"
        assert sent.claimed_at is None


def test_mri_registration_details_and_annual_serial():
    payload = token_payload()
    payload.pop("patient_phone")
    payload.update(age=42, unit="Unit A", mri_area="Brain", contrast=2, film=3,
                   report="Awaiting review", patient_source="walk_in")
    first = client.post("/api/v1/tokens", json=payload)
    assert first.status_code == 201, first.text
    record = first.json()
    assert record["serial_number"] == f"00001/{date.today().year % 100:02d}"
    assert record["token_date"] == date.today().isoformat()
    for key in ("age", "unit", "mri_area", "contrast", "film", "report", "patient_source"):
        assert record[key] == payload[key]
    second = client.post("/api/v1/tokens", json=payload)
    assert second.json()["serial_number"] == f"00002/{date.today().year % 100:02d}"
    payload["contrast"] = -1
    assert client.post("/api/v1/tokens", json=payload).status_code == 422


def test_registration_serial_is_global_and_resets_by_year():
    from app.database import SessionLocal
    from app.models import RegistrationCounter

    with SessionLocal() as db:
        db.add(RegistrationCounter(year=date.today().year - 1, sequence=999))
        db.add(Doctor(id="mri-two", name="MRI Two", department="MRI", designation="Radiographer",
                      room_number="206", waiting_room_id="waiting-room-1", token_prefix="MT"))
        db.commit()
    preview = client.get("/api/v1/registration-preview").json()
    assert preview["serial_number"] == f"00001/{date.today().year % 100:02d}"
    first = client.post("/api/v1/tokens", json=token_payload()).json()
    second = client.post("/api/v1/tokens", json=token_payload() | {"doctor_id": "mri-two"}).json()
    assert first["serial_number"].startswith("00001/")
    assert second["serial_number"].startswith("00002/")


def test_mri_details_are_not_exposed_on_waiting_room_display():
    from app.database import SessionLocal
    from app.service import QueueService

    created = client.post("/api/v1/tokens", json=token_payload() | {"report": "Private report", "mri_area": "Brain"})
    created.raise_for_status()
    with SessionLocal() as db:
        service = QueueService(db)
        display = service.display("WR-1")
        assert display.next_tokens[0].report is None
        assert display.next_tokens[0].mri_area is None
        service.call(created.json()["id"], "dr-khan")
        assert service.display("WR-1").current.report is None


def test_configured_designation_preserves_vip_without_assigning_room():
    created = client.post('/api/v1/lookups', json={
        'category': 'rank_relationship', 'value': 'brigadier_general', 'label': 'Brigadier General',
        'metadata_json': {'priority': 'vip'}, 'sort_order': 0,
    })
    assert created.status_code == 201, created.text
    source = client.post('/api/v1/lookups', json={'category': 'patient_source', 'value': 'opd', 'label': 'OPD'})
    assert source.status_code == 201
    payload = token_payload() | {'rank': 'brigadier_general', 'patient_source': 'opd', 'doctor_id': '',
                                 'room_number': '999', 'priority': 'normal'}
    response = client.post('/api/v1/tokens', json=payload)
    assert response.status_code == 201, response.text
    token = response.json()
    assert token['priority'] == 'vip'
    assert token['rank'] == 'brigadier_general'
    assert token['room_number'] == ''
    assert token['doctor_id'] is None
    assert client.patch(f"/api/v1/tokens/{token['id']}/room", json={'room_number': '999'}).status_code == 409
    assert client.get('/api/v1/dashboard/patients?status=vip').json()[0]['id'] == token['id']
    source_id = source.json()['id']
    client.patch(f'/api/v1/lookups/{source_id}', json={'is_active': False}).raise_for_status()
    assert client.post('/api/v1/tokens', json=payload).status_code == 422


def test_vips_sort_first_and_physical_call_is_silent_and_audited():
    normal = client.post('/api/v1/tokens', json=token_payload('Normal Patient')).json()
    vip = client.post('/api/v1/tokens', json=token_payload('VIP Patient') | {'priority': 'vip'}).json()
    assert client.get('/api/v1/tokens').json()[0]['id'] == vip['id']
    assert client.post(f"/api/v1/doctors/dr-khan/tokens/{normal['id']}/action", json={'action': 'call_physically'}).status_code == 409
    result = client.post(f"/api/v1/doctors/dr-khan/tokens/{vip['id']}/action", json={'action': 'call_physically'})
    assert result.status_code == 200, result.text
    assert result.json()['status'] == 'in_progress'
    assert result.json()['called_at'] and result.json()['started_at']
    from app.database import SessionLocal
    from app.models import AuditEvent, CallEvent
    with SessionLocal() as db:
        assert db.scalar(select(CallEvent).where(CallEvent.token_id == vip['id'])) is None
        assert db.scalar(select(AuditEvent).where(AuditEvent.token_id == vip['id'], AuditEvent.action == 'queue.call_physically'))
    display = client.get('/api/v1/displays/WR-1').json()
    assert display['current'] is None
    assert all(item['id'] != vip['id'] for item in display['next_tokens'])
    assert client.post(f"/api/v1/doctors/dr-khan/tokens/{vip['id']}/action", json={'action': 'cancel', 'reason': 'Patient requested'}).json()['status'] == 'cancelled'


def test_idle_radiographer_can_take_waiting_patient_but_cannot_take_vip_or_running_patient():
    from app.database import SessionLocal
    with SessionLocal() as db:
        db.add(Doctor(id='mri-idle', name='Idle Radiographer', department='MRI', designation='Radiographer',
                      room_number='209', waiting_room_id='waiting-room-1', token_prefix='MI'))
        db.commit()
    normal = client.post('/api/v1/tokens', json=token_payload()).json()
    vip = client.post('/api/v1/tokens', json=token_payload('VIP Patient') | {'priority': 'vip'}).json()
    available = client.get('/api/v1/doctors/mri-idle/available-patients')
    assert available.status_code == 200
    assert [item['id'] for item in available.json()] == [normal['id']]
    assert client.post(f"/api/v1/doctors/mri-idle/claim/{vip['id']}").status_code == 409
    claimed = client.post(f"/api/v1/doctors/mri-idle/claim/{normal['id']}")
    assert claimed.status_code == 200, claimed.text
    assert claimed.json()['doctor_id'] == 'mri-idle'
    assert claimed.json()['room_number'] == '209'
    assert claimed.json()['serial_number'] == normal['serial_number']
    assert client.get('/api/v1/doctors/mri-idle/available-patients').status_code == 409
    assert client.post(f"/api/v1/doctors/mri-idle/claim/{vip['id']}").status_code == 409
    client.post(f"/api/v1/doctors/dr-khan/tokens/{vip['id']}/action", json={'action': 'call_physically'}).raise_for_status()
    assert client.post(f"/api/v1/doctors/dr-khan/claim/{normal['id']}").status_code == 409


def test_called_patient_can_be_cancelled_and_disappears_from_display():
    token = client.post('/api/v1/tokens', json=token_payload()).json()
    client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").raise_for_status()
    result = client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={'action': 'cancel', 'reason': 'Patient left'})
    assert result.status_code == 200
    assert client.get('/api/v1/displays/WR-1').json()['current'] is None


def test_mri_report_columns_and_text_are_exported_safely():
    payload = token_payload() | {'report': '=1+1', 'age': 30, 'unit': 'Unit A', 'mri_area': 'Brain', 'contrast': 0, 'film': 2, 'patient_source': 'opd'}
    client.post('/api/v1/tokens', json=payload).raise_for_status()
    exported = client.get('/api/v1/reports/reception.xlsx')
    sheet = load_workbook(BytesIO(exported.content))['Reception Desk']
    assert [cell.value for cell in sheet[1]][:12] == ['Serial', 'Date', 'Service No./BA', 'Designation / Rank', 'Name', 'Age', 'Unit', 'MRI Area', 'Contrast', 'Film', 'Report', 'Patient Source']
    assert sheet['K2'].value == '=1+1' and sheet['K2'].data_type == 's'
    assert sheet['I2'].value == 0 and sheet['J2'].value == 2


def test_startup_seed_preserves_admin_dropdown_rooms_and_console_settings(monkeypatch):
    from app.database import SessionLocal
    from app.models import LookupOption
    from app.seed import seed
    monkeypatch.setenv('CMH_SMS_SEED_DEMO', 'false')
    seed()
    with SessionLocal() as db:
        rank = db.scalar(select(LookupOption).where(LookupOption.category == 'rank_relationship', LookupOption.value == 'brigadier_general'))
        rank.label, rank.is_active, rank.metadata_json = 'Custom General', False, {'priority': 'normal'}
        db.get(Doctor, 'dr-khan').room_number = '999'
        db.get(AppSetting, 'display').value = {'privacy_mode': 'full', 'next_token_count': 8}
        db.commit()
    seed()
    with SessionLocal() as db:
        rank = db.scalar(select(LookupOption).where(LookupOption.category == 'rank_relationship', LookupOption.value == 'brigadier_general'))
        assert rank.label == 'Custom General' and rank.is_active is False
        assert rank.metadata_json == {'priority': 'normal'}
        assert db.get(Doctor, 'dr-khan').room_number == '999'
        assert db.get(AppSetting, 'display').value['next_token_count'] == 8


def test_idle_claim_and_dashboard_details_respect_radiographer_scope():
    from app.database import SessionLocal
    with SessionLocal() as db:
        db.add(Doctor(id='mri-idle', name='Idle Radiographer', department='MRI', designation='Radiographer',
                      room_number='209', waiting_room_id='waiting-room-1', token_prefix='MI'))
        db.add(User(username='idle-radio', full_name='Idle Radiographer', password_hash=hash_password('Radiographer123!'),
                    role='radiographer', doctor_id='mri-idle', is_active=True))
        db.commit()
    normal = client.post('/api/v1/tokens', json=token_payload()).json()
    client.cookies.clear()
    client.post('/api/v1/auth/login', json={'username': 'idle-radio', 'password': 'Radiographer123!'}).raise_for_status()
    assert client.get('/api/v1/dashboard/patients').json()[0]['id'] == normal['id']
    assert client.get('/api/v1/doctors/dr-khan/available-patients').status_code == 403
    assert client.post(f"/api/v1/doctors/dr-khan/claim/{normal['id']}").status_code == 403
    assert client.get('/api/v1/doctors/mri-idle/available-patients').json()[0]['id'] == normal['id']
    result = client.post(f"/api/v1/doctors/mri-idle/claim/{normal['id']}")
    assert result.status_code == 200, result.text
    assert client.get('/api/v1/dashboard/patients?status=waiting').json()[0]['id'] == normal['id']


def test_reassignment_rejects_stale_source_queue():
    import pytest
    from fastapi import HTTPException

    from app.database import SessionLocal
    from app.service import QueueService
    with SessionLocal() as db:
        for index in (1, 2):
            db.add(Doctor(id=f'idle-{index}', name=f'Idle {index}', department='MRI', designation='Radiographer',
                          room_number=f'20{index}', waiting_room_id='waiting-room-1', token_prefix=f'I{index}'))
        db.commit()
    token = client.post('/api/v1/tokens', json=token_payload()).json()
    with SessionLocal() as stale, SessionLocal() as fresh:
        cached = stale.get(QueueToken, token['id'])
        assert cached.doctor_id == 'dr-khan'
        QueueService(fresh).transfer(token['id'], 'idle-1', 'First transfer', 'admin')
        with pytest.raises(HTTPException) as error:
            QueueService(stale).transfer(token['id'], 'idle-2', 'Stale claim', 'admin', expected_doctor_id='dr-khan')
        assert error.value.status_code == 409
    assert client.get('/api/v1/tokens').json()[0]['doctor_id'] == 'idle-1'


def test_shared_registration_has_no_room_and_any_radiographer_can_call_once():
    from app.database import SessionLocal
    with SessionLocal() as db:
        db.add(WaitingRoom(id='wr-other', code='WR-OTHER', name='Other Waiting Room', floor='1', display_label='Other'))
        db.add(Doctor(id='dr-other', name='Other Radiographer', department='MRI', designation='Radiographer',
                      room_number='301', waiting_room_id='wr-other', token_prefix='OT'))
        db.add(User(username='other-radio', full_name='Other Radiographer', password_hash=hash_password('Radiographer123!'),
                    role='radiographer', doctor_id='dr-other', is_active=True))
        db.commit()
    token = client.post('/api/v1/tokens', json={'patient_name': 'Shared Patient', 'patient_source': 'opd'}).json()
    assert token['doctor_id'] is None and token['room_number'] == '' and token['waiting_room'] == ''
    assert token['token_number'] == token['serial_number']
    assert client.get('/api/v1/displays/WR-1').json()['next_tokens'][0]['id'] == token['id']
    assert client.get('/api/v1/displays/WR-OTHER').json()['next_tokens'][0]['id'] == token['id']
    client.cookies.clear()
    client.post('/api/v1/auth/login', json={'username': 'other-radio', 'password': 'Radiographer123!'}).raise_for_status()
    assert client.get('/api/v1/tokens').json()[0]['id'] == token['id']
    called = client.post(f"/api/v1/doctors/dr-other/tokens/{token['id']}/call")
    assert called.status_code == 200, called.text
    assert called.json()['doctor_id'] == 'dr-other' and called.json()['room_number'] == '301'
    assert called.json()['serial_number'] == token['serial_number']
    client.cookies.clear()
    client.post('/api/v1/auth/login', json={'username': 'admin', 'password': 'AdminPass123!'}).raise_for_status()
    assert client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/call").status_code == 409
    assert client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={'action': 'start'}).status_code == 409
    another = client.post('/api/v1/tokens', json={'patient_name': 'Another Shared Patient'}).json()
    assert client.post(f"/api/v1/doctors/dr-other/tokens/{another['id']}/call").status_code == 409
    assert client.post(f"/api/v1/doctors/dr-other/tokens/{token['id']}/action", json={'action': 'cancel', 'reason': 'Patient left'}).status_code == 200
    assert client.post('/api/v1/doctors/dr-other/call-next').json()['id'] == another['id']


def test_radiographer_can_call_waiting_patient_from_other_waiting_room():
    from app.database import SessionLocal
    with SessionLocal() as db:
        db.add(Doctor(id='dr-other', name='Other Radiographer', department='MRI', designation='Radiographer',
                      room_number='301', waiting_room_id='waiting-room-1', token_prefix='OT'))
        db.commit()
    token = client.post('/api/v1/tokens', json=token_payload()).json()
    called = client.post(f"/api/v1/doctors/dr-other/tokens/{token['id']}/call")
    assert called.status_code == 200
    assert called.json()['doctor_id'] == 'dr-other' and called.json()['room_number'] == '301'
    assert called.json()['serial_number'] == token['serial_number']


def test_unassigned_vip_can_be_called_physically_by_any_radiographer():
    token = client.post('/api/v1/tokens', json={'patient_name': 'Shared VIP', 'priority': 'vip'}).json()
    assert token['doctor_id'] is None and not token['room_number']
    result = client.post(f"/api/v1/doctors/dr-khan/tokens/{token['id']}/action", json={'action': 'call_physically'})
    assert result.status_code == 200, result.text
    assert result.json()['status'] == 'in_progress' and result.json()['room_number'] == '205'
    assert client.get('/api/v1/displays/WR-1').json()['current'] is None


def test_long_mri_report_can_export_pdf_and_wait_stops_after_call():
    from app.database import SessionLocal
    from app.service import QueueService
    payload = token_payload() | {'report': 'Detailed MRI finding. ' * 400}
    token = client.post('/api/v1/tokens', json=payload).json()
    response = client.get('/api/v1/reports/reception.pdf')
    assert response.status_code == 200 and response.content.startswith(b'%PDF')
    with SessionLocal() as db:
        saved = db.get(QueueToken, token['id'])
        saved.created_at = datetime.utcnow() - timedelta(minutes=30)
        saved.called_at = saved.created_at + timedelta(minutes=5)
        saved.status = 'completed'
        db.commit()
        assert QueueService(db)._read(saved).waiting_minutes == 5


def install_summary_lookups():
    from app.database import SessionLocal
    from app.models import LookupOption
    from app.monthly_report import CLASSIFICATION_LOOKUPS
    with SessionLocal() as db:
        for category, options in CLASSIFICATION_LOOKUPS.items():
            for value, label in options:
                db.add(LookupOption(category=category, value=value, label=label, metadata_json={'report_code': value}))
        for group in ['officer', 'cadet', 'jco', 'or', 'nce']:
            db.add(LookupOption(category='rank_relationship', value=group, label=group, metadata_json={'report_group': group, 'priority': 'vip' if group == 'officer' else 'normal'}))
        db.commit()


def test_monthly_summary_classification_snapshot_exports_and_correction():
    from app.database import SessionLocal
    from app.models import AuditEvent, LookupOption
    install_summary_lookups()
    matrix = [
        ('self', 'serving', 'military', 'officer', 'serving_officer'),
        ('self', 'serving', 'military', 'cadet', 'serving_cadet'),
        ('self', 'serving', 'military', 'jco', 'serving_jco_or'),
        ('self', 'serving', 'military', 'or', 'serving_jco_or'),
        ('self', 'serving', 'military', 'nce', 'serving_nce'),
        ('self', '', 'civil', '', 'serving_civil'),
        ('self', 'retired', 'military', 'officer', 'retired_officer'),
        ('family', 'retired', 'military', 'officer', 'retired_officer'),
        ('self', 'retired', 'military', 'or', 'retired_or'),
        ('family', 'serving', 'military', 'officer', 'family_officer'),
        ('family', 'serving', 'military', 'jco', 'family_jco_or_nce'),
        ('family', 'serving', 'military', 'nce', 'family_jco_or_nce'),
        ('family', '', 'civil', '', 'family_civil'),
        ('self', '', 're', '', 're'),
        ('self', '', 'cne', '', 'cne'),
        ('family', 'retired', 'military', 'or', None),
    ]
    for person, status, entitlement, rank, expected in matrix:
        payload = {**token_payload(), 'doctor_id': None, 'waiting_room': '', 'rank': rank if person == 'self' else '', 'beneficiary_type': person, 'service_status': status, 'entitlement': entitlement, 'sponsor_rank': rank if person == 'family' else '', 'family_relationship': 'spouse' if person == 'family' else ''}
        response = client.post('/api/v1/tokens', json=payload)
        assert response.status_code == 201, response.text
        token = response.json()
        assert token['summary_category'] == expected
        assert token['doctor_id'] is None and token['room_number'] == ''
        if person == 'family':
            assert token['priority'] == 'normal'  # Sponsor VIP does not make the patient VIP.
    legacy = client.post('/api/v1/tokens', json={**token_payload(), 'rank': ''}).json()
    month = date.today().strftime('%Y-%m')
    query = {'month': month}
    data = client.get('/api/v1/reports/mri-summary', params=query).json()
    assert data['total'] == 17
    assert sum(data['totals'].values()) == 17
    assert data['totals']['unclassified'] == 2
    assert data['totals']['retired_officer'] == 2
    assert data['totals']['family_officer'] == 1
    with SessionLocal() as db:
        rank = db.scalar(select(LookupOption).where(LookupOption.category == 'rank_relationship', LookupOption.value == 'officer'))
        rank.metadata_json = {'report_group': 'or'}
        db.commit()
    assert client.get('/api/v1/reports/mri-summary', params=query).json() == data
    correction = client.patch(f"/api/v1/tokens/{legacy['id']}/classification", json={'beneficiary_type': 'family', 'entitlement': 'military', 'service_status': 'serving', 'sponsor_rank': 'jco', 'family_relationship': 'child', 'rank': ''})
    assert correction.status_code == 200, correction.text
    assert correction.json()['summary_category'] == 'family_jco_or_nce'
    with SessionLocal() as db:
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == 'token.classification_updated'))
    rows = client.get('/api/v1/reports/mri-summary/patients', params={**query, 'category': 'family_jco_or_nce'}).json()
    assert len(rows) == 3
    excel = client.get('/api/v1/reports/mri-summary.xlsx', params=query)
    assert excel.status_code == 200
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet['B6'].value == 'SERVING'
    assert sheet.cell(sheet.max_row, sheet.max_column).value == 17
    assert sheet.cell(sheet.max_row, sheet.max_column-1).value == 1
    pdf = client.get('/api/v1/reports/mri-summary.pdf', params=query)
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF')
    client.cookies.clear()
    assert client.get('/api/v1/reports/mri-summary', params=query).status_code == 401


def test_summary_completion_dates_use_dhaka_and_invalid_inputs():
    from app.database import SessionLocal
    install_summary_lookups()
    ids = [client.post('/api/v1/tokens', json={**token_payload(), 'rank': ''}).json()['id'] for _ in range(4)]
    with SessionLocal() as db:
        for identity, completed, status in zip(ids, [datetime(2026, 5, 31, 17, 59), datetime(2026, 5, 31, 18), datetime(2026, 6, 30, 17, 59), datetime(2026, 6, 30, 18)], ['completed']*4):
            token = db.get(QueueToken, identity)
            token.completed_at = completed
            token.status = status
            token.summary_category = 're'
        db.commit()
    report = client.get('/api/v1/reports/mri-summary', params={'month': '2026-06', 'basis': 'completed'}).json()
    assert report['total'] == 2
    assert len(report['rows']) == 30
    assert report['rows'][0]['total'] == 1 and report['rows'][-1]['total'] == 1
    assert client.get('/api/v1/reports/mri-summary', params={'month': '2026-13'}).status_code == 422
    assert client.get('/api/v1/reports/mri-summary', params={'month': '2026-06', 'basis': 'anything'}).status_code == 422
    invalid = client.post('/api/v1/tokens', json={**token_payload(), 'rank': '', 'beneficiary_type': 'family', 'entitlement': 'military', 'family_relationship': 'spouse'})
    assert invalid.status_code == 422
