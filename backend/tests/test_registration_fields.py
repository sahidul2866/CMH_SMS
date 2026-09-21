from io import BytesIO
from openpyxl import load_workbook
from test_queue import client, token_payload, setup_function  # noqa: F401


def field(key='custom_referral', kind='text', **changes):
    return {'key': key, 'label': 'Referral reference', 'type': kind, 'enabled': True,
            'required': False, 'options': [], **changes}


def save_policy(**changes):
    config = client.get('/api/v1/registration-fields').json()
    config.update(changes)
    response = client.put('/api/v1/registration-fields', json={'value': config})
    assert response.status_code == 200, response.text
    return response.json()


def test_disabled_inputs_are_rejected_and_history_is_preserved():
    original = {**token_payload(), 'unit': 'Original unit', 'age': 42}
    saved = client.post('/api/v1/tokens', json=original).json()
    config = client.get('/api/v1/registration-fields').json()
    save_policy(enabled=[key for key in config['enabled'] if key not in ['unit', 'age', 'priority']])
    assert client.post('/api/v1/tokens', json=original).status_code == 422
    assert client.patch(f"/api/v1/tokens/{saved['id']}", json={**original, 'unit': 'replacement'}).status_code == 422
    allowed = {key: value for key, value in token_payload('Changed name').items() if key != 'priority'}
    response = client.patch(f"/api/v1/tokens/{saved['id']}", json=allowed)
    assert response.status_code == 200, response.text
    assert response.json()['unit'] == 'Original unit'
    assert response.json()['age'] == 42
    created = client.post('/api/v1/tokens', json=allowed)
    assert created.status_code == 201, created.text
    assert created.json()['unit'] is None
    assert created.json()['priority'] == 'normal'
    save_policy(enabled=config['enabled'])
    assert client.patch(f"/api/v1/tokens/{saved['id']}", json={**original, 'unit': 'Updated unit'}).json()['unit'] == 'Updated unit'


def test_invalid_settings_and_legacy_required_only_settings():
    from app.registration import FIELDS
    response = client.put('/api/v1/registration-fields', json={'value': {'required': ['patient_name', 'age']}})
    assert response.status_code == 200
    assert set(response.json()['enabled']) == set(FIELDS)
    for changes in [
        {'enabled': ['age'], 'required': ['age']},
        {'enabled': ['patient_name'], 'required': ['patient_name', 'age']},
        {'enabled': ['patient_name', 'unknown'], 'required': ['patient_name']},
        {'required': ['patient_name'], 'custom': [field(enabled=False, required=True)]},
        {'required': ['patient_name'], 'custom': [field(kind='select')]},
        {'required': ['patient_name'], 'custom': [field(), field()]},
    ]:
        result = client.put('/api/v1/registration-fields', json={'value': changes})
        assert result.status_code == 422, result.text


def test_custom_values_validation_history_and_excel():
    fields = [field(required=True), field('custom_score', 'number', label='Score', required=True),
              field('custom_date', 'date', label='Referral date'),
              field('custom_side', 'select', label='Side', options=['Left', 'Right'])]
    save_policy(custom=fields)
    assert client.post('/api/v1/tokens', json=token_payload()).status_code == 422
    values = {'custom_referral': '=NotAFormula', 'custom_score': 0, 'custom_date': '2026-09-21', 'custom_side': 'Left'}
    response = client.post('/api/v1/tokens', json={**token_payload(), 'custom_fields': values})
    assert response.status_code == 201, response.text
    token = response.json()
    assert token['custom_fields'] == values
    for key, value in [('custom_score', '12'), ('custom_score', True), ('custom_score', 10 ** 400), ('custom_date', '2026-02-31'), ('custom_side', 'Other'), ('custom_referral', 5), ('custom_unknown', 'x')]:
        result = client.post('/api/v1/tokens', json={**token_payload(), 'custom_fields': {**values, key: value}})
        assert result.status_code == 422, result.text
    fields[0].update(enabled=False, required=False)
    save_policy(custom=fields)
    assert client.patch(f"/api/v1/tokens/{token['id']}", json={**token_payload(), 'custom_fields': values}).status_code == 422
    active_values = {key: value for key, value in values.items() if key != 'custom_referral'}
    updated = client.patch(f"/api/v1/tokens/{token['id']}", json={**token_payload('Updated patient'), 'custom_fields': active_values})
    assert updated.status_code == 200, updated.text
    assert updated.json()['custom_fields']['custom_referral'] == '=NotAFormula'
    exported = client.get('/api/v1/reports/reception.xlsx')
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(BytesIO(exported.content)).active
    column = [cell.value for cell in sheet[1]].index('Referral reference') + 1
    assert sheet.cell(2, column).value == '=NotAFormula'
    assert sheet.cell(2, column).data_type == 's'
    fresh = client.post('/api/v1/tokens', json={**token_payload(), 'custom_fields': active_values})
    assert 'custom_referral' not in fresh.json()['custom_fields']


def test_definition_edits_cannot_destroy_history():
    definition = field()
    save_policy(custom=[definition])
    for custom in [[], [{**definition, 'type': 'number'}], [{**definition, 'key': 'custom_different'}]]:
        result = client.put('/api/v1/registration-fields', json={'value': {'required': ['patient_name'], 'custom': custom}})
        assert result.status_code == 422
    response = save_policy(custom=[{**definition, 'label': 'Updated display label'}])
    assert response['custom'][0]['key'] == definition['key']


def test_disabled_family_fields_do_not_block_registration():
    from app.seed import seed
    seed()
    config = client.get('/api/v1/registration-fields').json()
    save_policy(required=['patient_name'], enabled=[key for key in config['enabled'] if key not in ['sponsor_rank', 'family_relationship', 'rank']])
    payload = {**token_payload(), 'beneficiary_type': 'family', 'entitlement': 'military', 'service_status': 'serving'}
    payload.pop('rank')
    result = client.post('/api/v1/tokens', json=payload)
    assert result.status_code == 201, result.text
    assert result.json()['summary_category'] is None
    assert client.patch(f"/api/v1/tokens/{result.json()['id']}/classification", json={'sponsor_rank': 'captain'}).status_code == 422


def test_configuration_is_audited_and_survives_seed():
    from app.seed import seed
    definition = field()
    before = save_policy(custom=[definition])
    seed()
    assert client.get('/api/v1/registration-fields').json() == before
    events = client.get('/api/v1/audit').json()
    assert any(event['action'] == 'registration_fields.updated' and event['detail']['custom'] == [definition] for event in events)


def test_direct_classification_preserves_inputs_and_survives_patient_edits():
    payload = {**token_payload(), 'age': 35, 'unit': 'Test unit'}
    token = client.post('/api/v1/tokens', json=payload).json()
    result = client.patch(f"/api/v1/tokens/{token['id']}/classification", json={'mode': 'direct', 'summary_category': 'family_officer'})
    assert result.status_code == 200, result.text
    assert result.json()['summary_category_source'] == 'manual'
    for key in ['rank', 'unit', 'age', 'beneficiary_type', 'entitlement']:
        assert result.json()[key] == token[key]
    edited = client.patch(f"/api/v1/tokens/{token['id']}", json={**payload, 'unit': 'Changed unit'})
    assert edited.status_code == 200, edited.text
    assert edited.json()['summary_category'] == 'family_officer'
    assert edited.json()['summary_category_source'] == 'manual'
    from app.seed import seed
    seed()
    automatic = client.patch(f"/api/v1/tokens/{token['id']}/classification", json={'mode': 'inputs', 'entitlement': 're'})
    assert automatic.status_code == 200, automatic.text
    assert automatic.json()['summary_category'] == 're'
    assert automatic.json()['summary_category_source'] == 'automatic'
    assert automatic.json()['rank'] == token['rank']
    events = client.get('/api/v1/audit').json()
    event = next(item for item in events if item['action'] == 'token.classification_updated' and item['detail']['mode'] == 'direct')
    assert event['actor'] == 'admin'
    assert event['detail']['before']['summary_category'] is None
    assert event['detail']['after']['summary_category'] == 'family_officer'


def test_direct_classification_requires_supported_explicit_category_only():
    token = client.post('/api/v1/tokens', json=token_payload()).json()
    path = f"/api/v1/tokens/{token['id']}/classification"
    for payload in [{'mode': 'direct'}, {'mode': 'direct', 'summary_category': 'invented'},
                    {'mode': 'direct', 'summary_category': 're', 'rank': 'captain'},
                    {'mode': 'inputs', 'summary_category': 're'},
                    {'mode': 'direct', 'summary_category': 're', 'patient_name': 'Overwrite'}]:
        assert client.patch(path, json=payload).status_code == 422
    assert client.patch(path, json={'mode': 'direct', 'summary_category': None}).status_code == 200


def test_direct_classification_works_with_disabled_inputs_and_required_registration_fields():
    token = client.post('/api/v1/tokens', json=token_payload()).json()
    save_policy(enabled=['patient_name', 'age'], required=['patient_name', 'age'])
    result = client.patch(f"/api/v1/tokens/{token['id']}/classification", json={'mode': 'direct', 'summary_category': 're'})
    assert result.status_code == 200, result.text
    assert result.json()['age'] is None
    assert result.json()['summary_category'] == 're'
    report = client.get('/api/v1/reports/mri-summary', params={'month': token['token_date'][:7]}).json()
    assert report['totals']['re'] == 1
    assert client.patch(f"/api/v1/tokens/{token['id']}/classification", json={'mode': 'inputs', 'entitlement': 're'}).status_code == 422


def test_classification_metadata_uses_classify_permission_without_settings_permission():
    from app.auth import hash_password
    from app.database import SessionLocal
    from app.models import RoleDefinition, User
    from app.seed import seed
    seed()
    with SessionLocal() as db:
        db.add(RoleDefinition(name='classify_only', display_name='Classification only', access_profile='reception', permissions=['queue.serial.create']))
        db.add(User(username='classifier', full_name='Test Classifier', role='classify_only', password_hash=hash_password('TestClassify123!'), is_active=True))
        db.commit()
    client.cookies.clear()
    client.post('/api/v1/auth/login', json={'username': 'classifier', 'password': 'TestClassify123!'}).raise_for_status()
    assert client.get('/api/v1/mri-summary-mapping').status_code == 403
    assert client.get('/api/v1/lookups').status_code == 403
    options = client.get('/api/v1/classification-options')
    assert options.status_code == 200, options.text
    assert len(options.json()['columns']) == 12
    assert options.json()['lookups']
    assert options.json()['registration']['enabled']
    client.cookies.clear()
    assert client.get('/api/v1/classification-options').status_code == 401
