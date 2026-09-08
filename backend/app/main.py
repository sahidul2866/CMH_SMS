from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import secrets
import threading
import time
from contextlib import asynccontextmanager
from datetime import date, datetime
from html import escape
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from fastapi import (
    Depends,
    FastAPI,
    Header,
    HTTPException,
    Query,
    Request,
    Response,
    WebSocket,
    WebSocketDisconnect,
    status,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from sqlalchemy import delete, func, or_, select, text
from sqlalchemy.orm import Session

from .announcement import announcement_engine
from .appointment_service import AppointmentService
from .auth import (
    COOKIE_SECURE,
    DUMMY_PASSWORD_HASH,
    SESSION_COOKIE,
    SESSION_HOURS,
    current_user,
    has_permission,
    hash_password,
    new_session,
    require_permission,
    revoke_session,
    user_for_session,
    verify_password,
    websocket_user,
)
from .database import DATABASE_URL, SessionLocal, get_db
from .models import (
    Appointment,
    AppSetting,
    AuditEvent,
    DeviceEndpoint,
    Doctor,
    Holiday,
    LookupOption,
    QueueControl,
    QueueToken,
    RoleDefinition,
    SmsMessage,
    User,
    UserSession,
    WaitingRoom,
)
from .monthly_report import router as monthly_report_router
from .notification import SmsService, sms_outbox_worker
from .realtime import waiting_room_hub
from .schemas import (
    AppointmentCreate,
    AppointmentRead,
    AppointmentUpdate,
    AuditRead,
    DashboardDoctorRead,
    DashboardRead,
    DashboardRoomRead,
    DeviceCreate,
    DeviceRead,
    DisplayRead,
    DoctorCreate,
    DoctorUpdate,
    HolidayCreate,
    HolidayRead,
    LoginRequest,
    LookupCreate,
    LookupRead,
    LookupUpdate,
    PasswordChange,
    PasswordReset,
    PatientCreate,
    PatientRead,
    PriorityUpdateRequest,
    QueueActionRequest,
    QueueControlUpdate,
    RoleCreate,
    RoleRead,
    RoleUpdate,
    SettingUpdate,
    SlotCreate,
    SlotRead,
    TokenCreate,
    TokenRead,
    TokenRoomUpdateRequest,
    TransferRequest,
    UserCreate,
    UserRead,
    UserUpdate,
    WaitingRoomCreate,
)
from .service import QueueService


@asynccontextmanager
async def lifespan(_: FastAPI):
    validate_production_config()
    with SessionLocal() as db:
        db.execute(delete(UserSession).where(UserSession.expires_at <= datetime.utcnow()))
        db.commit()
    await announcement_engine.start()
    await sms_outbox_worker.start()
    yield
    await sms_outbox_worker.stop()
    await announcement_engine.stop()


ENVIRONMENT = os.getenv("CMH_SMS_ENVIRONMENT", "development").lower()
logger = logging.getLogger("uvicorn.error")
login_attempts: dict[str, list[float]] = {}
login_attempts_lock = threading.Lock()


def serialize_user(user: User, db: Session) -> dict:
    role = db.get(RoleDefinition, user.role)
    return {
        "id": user.id, "username": user.username, "full_name": user.full_name,
        "role": user.role, "access_profile": role.access_profile if role else user.role,
        "permissions": role.permissions if role else ([] if user.role != "admin" else ["*"]),
        "doctor_id": user.doctor_id, "is_active": user.is_active,
    }


def user_access_profile(user: User, db: Session) -> str:
    role = db.get(RoleDefinition, user.role)
    return role.access_profile if role else user.role


def validate_production_config() -> None:
    if ENVIRONMENT != "production":
        return
    problems: list[str] = []
    if DATABASE_URL.startswith("sqlite"):
        problems.append("CMH_SMS_DATABASE_URL must use PostgreSQL")
    if not COOKIE_SECURE:
        problems.append("CMH_SMS_COOKIE_SECURE must be true")
    if os.getenv("CMH_SMS_PUBLIC_HTTPS", "false").lower() not in {"1", "true", "yes"}:
        problems.append("CMH_SMS_PUBLIC_HTTPS must be true")
    if os.getenv("CMH_SMS_ALLOWED_HOSTS", "*").strip() in {"", "*"}:
        problems.append("CMH_SMS_ALLOWED_HOSTS must list the approved server hostname/IP")
    if problems:
        raise RuntimeError("Unsafe production configuration: " + "; ".join(problems))


def enforce_login_rate_limit(client: str) -> None:
    now = time.monotonic()
    with login_attempts_lock:
        recent = [attempt for attempt in login_attempts.get(client, []) if now - attempt < 300]
        if len(recent) >= 10:
            raise HTTPException(429, "Too many login attempts. Try again later.")
        recent.append(now)
        login_attempts[client] = recent


def clear_login_rate_limit(client: str) -> None:
    with login_attempts_lock:
        login_attempts.pop(client, None)


app = FastAPI(
    title="CMH Smart Serial API",
    version="1.0.0",
    lifespan=lifespan,
    docs_url="/docs" if ENVIRONMENT != "production" else None,
    redoc_url=None,
)

def _normalize_origin(value: str) -> str:
    if not value:
        return ""
    cleaned = value.strip().rstrip("/")
    if not cleaned:
        return ""
    if "://" not in cleaned:
        return "https://" + cleaned.lstrip(".")
    return cleaned


def _collect_cors_origins() -> list[str]:
    origins: list[str] = []
    for env_name in (
        "CMH_SMS_CORS_ORIGINS",
        "CMH_SMS_FRONTEND_ORIGIN",
        "CMH_SMS_DEV_TUNNEL_URL",
        "CMH_SMS_TUNNEL_URL",
    ):
        for item in os.getenv(env_name, "").split(","):
            normalized = _normalize_origin(item)
            if normalized:
                origins.append(normalized)
    seen: set[str] = set()
    unique: list[str] = []
    for origin in origins:
        if origin not in seen:
            seen.add(origin)
            unique.append(origin)
    return unique


allowed_hosts = [item.strip() for item in os.getenv("CMH_SMS_ALLOWED_HOSTS", "*").split(",") if item.strip()]
allowed_hosts = allowed_hosts or ["*"]


def _expand_allowed_hosts_from_origins(origins: list[str]) -> list[str]:
    expanded: list[str] = []
    for origin in origins:
        parsed = urlparse(origin)
        host = parsed.hostname or ""
        if host:
            expanded.append(host)
    return expanded


cors_origins = _collect_cors_origins()
allow_origin_regex = r"https?://.*\.devtunnels\.ms"
allowed_hosts = list(dict.fromkeys(allowed_hosts + _expand_allowed_hosts_from_origins(cors_origins)))
app.add_middleware(TrustedHostMiddleware, allowed_hosts=allowed_hosts or ["*"])
if cors_origins or allow_origin_regex:
    app.add_middleware(
        CORSMiddleware,
        allow_origins=cors_origins,
        allow_origin_regex=allow_origin_regex,
        allow_credentials=True,
        allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE"],
        allow_headers=["Content-Type", "X-Device-Key"],
    )


@app.middleware("http")
async def production_headers_and_request_log(request: Request, call_next):
    request_id = request.headers.get("X-Request-ID", "")[:80] or secrets.token_hex(12)
    started = time.perf_counter()
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data:; font-src 'self'; connect-src 'self' ws: wss:; "
        "object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
    )
    content_type = response.headers.get("content-type", "")
    if request.url.path.startswith("/api/") or content_type.startswith("text/html"):
        response.headers["Cache-Control"] = "no-store"
    elif re.search(r"-[A-Z0-9]{8,}\.(?:js|css)$", request.url.path):
        response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    else:
        response.headers["Cache-Control"] = "no-cache"
    if request.url.scheme == "https" or os.getenv("CMH_SMS_PUBLIC_HTTPS", "false").lower() in {"1", "true", "yes"}:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    logger.info(
        json.dumps(
            {
                "event": "http_request",
                "request_id": request_id,
                "method": request.method,
                "path": request.url.path,
                "status": response.status_code,
                "duration_ms": round((time.perf_counter() - started) * 1000, 2),
            },
            separators=(",", ":"),
        )
    )
    return response


@app.get("/api/v1/health")
def health():
    return {"status": "ok", "service": "cmh-smart-serial", "audio": announcement_engine.status()}


@app.get("/api/v1/ready")
def ready(db: Session = Depends(get_db)):
    try:
        db.execute(text("SELECT 1"))
    except Exception as exc:
        logger.exception("Readiness database check failed")
        raise HTTPException(503, "Database unavailable") from exc
    return {"status": "ready", "database": "ok", "environment": ENVIRONMENT}


def queue_announcement(db: Session, token: TokenRead) -> None:
    setting = db.get(AppSetting, "announcement")
    announcement_engine.enqueue(token, setting.value if setting else {})


def enforce_doctor_access(user: User, doctor_id: str, db: Session) -> None:
    if user_access_profile(user, db) == "radiographer" and user.doctor_id != doctor_id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This account is assigned to another doctor")


def validated_setting(key: str, value: dict) -> dict:
    allowed = {
        "display": {"privacy_mode", "next_token_count", "ticker_message"},
        "queue": {"ordering_policy", "recall_limit", "late_grace_minutes"},
        "announcement": {"language_order", "repeat_count", "rate", "volume", "voice_mode", "cache_max_files"},
    }
    if key not in allowed:
        raise HTTPException(404, "Setting not found")
    unknown = set(value) - allowed[key]
    if unknown:
        raise HTTPException(422, f"Unknown {key} setting: {sorted(unknown)[0]}")

    def integer(name: str, minimum: int, maximum: int) -> None:
        item = value.get(name)
        if isinstance(item, bool) or not isinstance(item, int) or not minimum <= item <= maximum:
            raise HTTPException(422, f"{name} must be an integer from {minimum} to {maximum}")

    def number(name: str, minimum: float, maximum: float) -> None:
        item = value.get(name)
        if isinstance(item, bool) or not isinstance(item, (int, float)) or not minimum <= item <= maximum:
            raise HTTPException(422, f"{name} must be a number from {minimum} to {maximum}")

    if key == "display":
        if value.get("privacy_mode") not in {"initials", "token_only", "full"}:
            raise HTTPException(422, "privacy_mode is invalid")
        integer("next_token_count", 1, 10)
        if not isinstance(value.get("ticker_message"), str) or len(value["ticker_message"]) > 240:
            raise HTTPException(422, "ticker_message must contain at most 240 characters")
    elif key == "queue":
        if value.get("ordering_policy") not in {"priority_then_sequence", "sequence"}:
            raise HTTPException(422, "ordering_policy is invalid")
        integer("recall_limit", 0, 10)
        integer("late_grace_minutes", 0, 240)
    else:
        languages = value.get("language_order")
        if not isinstance(languages, list) or not languages or any(item not in {"bn", "en"} for item in languages):
            raise HTTPException(422, "language_order must contain bn and/or en")
        integer("repeat_count", 1, 3)
        number("rate", 0.5, 1.5)
        number("volume", 0, 1)
        if value.get("voice_mode", "auto") not in {"auto", "offline_neural", "offline"}:
            raise HTTPException(422, "voice_mode must be auto, offline_neural, or offline")
        cache_max_files = value.get("cache_max_files", 40)
        if isinstance(cache_max_files, bool) or not isinstance(cache_max_files, int) or not 10 <= cache_max_files <= 500:
            raise HTTPException(422, "cache_max_files must be an integer from 10 to 500")
    return value


@app.post("/api/v1/auth/login", response_model=UserRead)
def login(payload: LoginRequest, request: Request, response: Response, db: Session = Depends(get_db)):
    client = request.client.host if request.client else "unknown"
    enforce_login_rate_limit(client)
    user = db.scalar(select(User).where(User.username == payload.username.lower()))
    password_valid = verify_password(payload.password, user.password_hash if user else DUMMY_PASSWORD_HASH)
    if not user or not user.is_active or not password_valid:
        db.add(
            AuditEvent(
                action="auth.login_failed",
                actor=payload.username.lower(),
                detail={"client": client},
            )
        )
        db.commit()
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid username or password")
    raw = new_session(db, user)
    clear_login_rate_limit(client)
    db.add(
        AuditEvent(
            action="auth.login",
            actor=user.username,
            detail={"client": client},
        )
    )
    db.commit()
    response.set_cookie(
        SESSION_COOKIE,
        raw,
        httponly=True,
        secure=COOKIE_SECURE,
        samesite="strict",
        max_age=SESSION_HOURS * 60 * 60,
        path="/",
    )
    return serialize_user(user, db)


@app.post("/api/v1/auth/logout", status_code=204)
def logout(request: Request, response: Response, user: User = Depends(current_user), db: Session = Depends(get_db)):
    revoke_session(db, request.cookies.get(SESSION_COOKIE))
    db.add(AuditEvent(action="auth.logout", actor=user.username, detail={}))
    db.commit()
    response.delete_cookie(SESSION_COOKIE, path="/", secure=COOKIE_SECURE, httponly=True, samesite="strict")


@app.get("/api/v1/auth/me", response_model=UserRead)
def me(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return serialize_user(user, db)


@app.get("/api/v1/auth/session", response_model=UserRead | None)
def auth_session(request: Request, db: Session = Depends(get_db)):
    user = user_for_session(db, request.cookies.get(SESSION_COOKIE))
    return serialize_user(user, db) if user else None


@app.post("/api/v1/auth/password", status_code=204)
def change_password(
    payload: PasswordChange,
    request: Request,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(400, "Current password is incorrect")
    if payload.current_password == payload.new_password:
        raise HTTPException(400, "New password must be different")
    user.password_hash = hash_password(payload.new_password)
    current_session_id = hashlib.sha256((request.cookies.get(SESSION_COOKIE) or "").encode()).hexdigest()
    db.execute(
        delete(UserSession).where(
            UserSession.user_id == user.id,
            UserSession.id != current_session_id,
        )
    )
    db.add(AuditEvent(action="auth.password_changed", actor=user.username, detail={}))
    db.commit()


PERMISSION_CATALOG = [
    {"key": "pages.reception", "label": "Open Reception Desk", "group": "Pages"},
    {"key": "pages.radiographer", "label": "Open Radiographers", "group": "Pages"},
    {"key": "pages.display", "label": "Open Waiting Display", "group": "Pages"},
    {"key": "pages.reports", "label": "Open Reports & Audit", "group": "Pages"},
    {"key": "pages.dashboard", "label": "Open Radiography Dashboard", "group": "Pages"},
    {"key": "queue.serial.create", "label": "Generate patient serials", "group": "Queue"},
    {"key": "queue.view", "label": "View permitted queue", "group": "Queue"},
    {"key": "queue.call", "label": "Call patients", "group": "Queue"},
    {"key": "queue.action", "label": "Start, complete, skip and recall", "group": "Queue"},
    {"key": "queue.priority", "label": "Change patient priority", "group": "Queue"},
    {"key": "queue.transfer", "label": "Transfer patients", "group": "Queue"},
    {"key": "queue.print", "label": "Print and reprint serial slips", "group": "Queue"},
    {"key": "queue.pause", "label": "Pause and resume a queue", "group": "Queue"},
    {"key": "audio.announce", "label": "Play patient audio announcements", "group": "Audio"},
    {"key": "directory.view", "label": "View radiographers and waiting rooms", "group": "Directory"},
    {"key": "directory.manage", "label": "Create radiographers and waiting rooms", "group": "Directory"},
    {"key": "doctor.room.manage", "label": "Change radiographer room numbers", "group": "Directory"},
    {"key": "master_data.view", "label": "View active dropdown values", "group": "Master data"},
    {"key": "display.view", "label": "View waiting-room display data", "group": "Display"},
    {"key": "reports.view", "label": "View reports and audit", "group": "Reports"},
    {"key": "dashboard.view", "label": "View role-scoped radiography dashboard", "group": "Reports"},
    {"key": "settings.manage", "label": "Manage queue and display settings", "group": "Settings"},
    {"key": "master_data.manage", "label": "Manage dropdown values", "group": "Settings"},
    {"key": "holidays.view", "label": "View holidays", "group": "Operations"},
    {"key": "holidays.manage", "label": "Manage holidays", "group": "Operations"},
    {"key": "devices.view", "label": "View registered devices", "group": "Operations"},
    {"key": "devices.manage", "label": "Register devices", "group": "Operations"},
    {"key": "sms.view", "label": "View SMS outbox", "group": "Operations"},
    {"key": "sms.retry", "label": "Retry failed SMS messages", "group": "Operations"},
    {"key": "users.view", "label": "View users", "group": "Users"},
    {"key": "users.create", "label": "Create users", "group": "Users"},
    {"key": "users.update", "label": "Update user details", "group": "Users"},
    {"key": "users.delete", "label": "Delete users", "group": "Users"},
    {"key": "users.password.reset", "label": "Reset other users' passwords", "group": "Users"},
    {"key": "users.roles.assign", "label": "Assign roles to users", "group": "Users"},
    {"key": "roles.view", "label": "View roles and permissions", "group": "Roles"},
    {"key": "roles.manage", "label": "Create, edit and delete roles", "group": "Roles"},
]
VALID_PERMISSIONS = {item["key"] for item in PERMISSION_CATALOG}


@app.get("/api/v1/permissions")
def permissions(_: User = Depends(require_permission("roles.view"))):
    return PERMISSION_CATALOG


@app.get("/api/v1/users", response_model=list[UserRead])
def users(_: User = Depends(require_permission("users.view")), db: Session = Depends(get_db)):
    return [serialize_user(user, db) for user in db.scalars(select(User).order_by(User.username))]


@app.get("/api/v1/roles", response_model=list[RoleRead])
def roles(_: User = Depends(require_permission("roles.view")), db: Session = Depends(get_db)):
    return list(db.scalars(select(RoleDefinition).order_by(RoleDefinition.display_name)))


@app.post("/api/v1/roles", response_model=RoleRead, status_code=201)
def create_role(payload: RoleCreate, _: User = Depends(require_permission("roles.manage")), db: Session = Depends(get_db)):
    if db.get(RoleDefinition, payload.name):
        raise HTTPException(409, "Role already exists")
    values = payload.model_dump()
    values["permissions"] = sorted(set(values["permissions"]) & VALID_PERMISSIONS)
    role = RoleDefinition(**values, is_system=False)
    db.add(role)
    db.commit()
    db.refresh(role)
    return role


@app.patch("/api/v1/roles/{role_name}", response_model=RoleRead)
def update_role(role_name: str, payload: RoleUpdate, _: User = Depends(require_permission("roles.manage")), db: Session = Depends(get_db)):
    role = db.get(RoleDefinition, role_name)
    if not role:
        raise HTTPException(404, "Role not found")
    values = payload.model_dump(exclude_unset=True)
    if role.name == "admin" and values.get("access_profile", "admin") != "admin":
        raise HTTPException(409, "The built-in administrator role must remain unrestricted")
    if "permissions" in values:
        values["permissions"] = sorted(set(values["permissions"] or []) & VALID_PERMISSIONS)
    for key, value in values.items():
        setattr(role, key, value)
    db.commit()
    db.refresh(role)
    return role


@app.delete("/api/v1/roles/{role_name}", status_code=204)
def delete_role(role_name: str, _: User = Depends(require_permission("roles.manage")), db: Session = Depends(get_db)):
    role = db.get(RoleDefinition, role_name)
    if not role:
        raise HTTPException(404, "Role not found")
    if role.is_system:
        raise HTTPException(409, "System roles cannot be deleted")
    if db.scalar(select(func.count()).select_from(User).where(User.role == role_name)):
        raise HTTPException(409, "Reassign users before deleting this role")
    db.delete(role)
    db.commit()


@app.post("/api/v1/users", response_model=UserRead, status_code=201)
def create_user(payload: UserCreate, _: User = Depends(require_permission("users.create")), db: Session = Depends(get_db)):
    username = payload.username.lower()
    if db.scalar(select(User).where(User.username == username)):
        raise HTTPException(409, "Username already exists")
    role = db.get(RoleDefinition, payload.role)
    if not role:
        raise HTTPException(422, "Assigned role does not exist")
    if role.access_profile == "radiographer" and not payload.doctor_id:
        raise HTTPException(422, "Radiographer accounts must be assigned to a doctor")
    if payload.doctor_id and not db.get(Doctor, payload.doctor_id):
        raise HTTPException(422, "Assigned doctor does not exist")
    user = User(
        username=username,
        full_name=payload.full_name,
        password_hash=hash_password(payload.password),
        role=payload.role,
        doctor_id=payload.doctor_id,
        is_active=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return serialize_user(user, db)


@app.patch("/api/v1/users/{user_id}", response_model=UserRead)
def update_user(
    user_id: str,
    payload: UserUpdate,
    administrator: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, "User not found")
    if user.id == administrator.id and payload.is_active is False:
        raise HTTPException(409, "You cannot disable your own account")
    values = payload.model_dump(exclude_unset=True)
    password = values.pop("password", None)
    if password:
        if not has_permission(administrator, db, "users.password.reset"):
            raise HTTPException(403, "Password reset permission required")
        user.password_hash = hash_password(password)
    assignment_fields = {"role", "doctor_id"} & values.keys()
    general_fields = set(values) - assignment_fields
    if assignment_fields and not has_permission(administrator, db, "users.roles.assign"):
        raise HTTPException(403, "Role assignment permission required")
    if general_fields and not has_permission(administrator, db, "users.update"):
        raise HTTPException(403, "User update permission required")
    for key, value in values.items():
        setattr(user, key, value)
    definition = db.get(RoleDefinition, user.role)
    if not definition:
        raise HTTPException(422, "Assigned role does not exist")
    if definition.access_profile == "radiographer" and not user.doctor_id:
        raise HTTPException(422, "Radiographer accounts must be assigned to a doctor")
    if user.doctor_id and not db.get(Doctor, user.doctor_id):
        raise HTTPException(422, "Assigned doctor does not exist")
    db.commit()
    db.refresh(user)
    return serialize_user(user, db)


@app.post("/api/v1/users/{user_id}/password", status_code=204)
def reset_user_password(
    user_id: str,
    payload: PasswordReset,
    administrator: User = Depends(require_permission("users.password.reset")),
    db: Session = Depends(get_db),
):
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, "User not found")
    user.password_hash = hash_password(payload.new_password)
    db.execute(delete(UserSession).where(UserSession.user_id == user.id))
    db.add(AuditEvent(action="auth.password_reset", actor=administrator.username, detail={"user": user.username}))
    db.commit()


@app.delete("/api/v1/users/{user_id}", status_code=204)
def delete_user(user_id: str, administrator: User = Depends(require_permission("users.delete")), db: Session = Depends(get_db)):
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, "User not found")
    if user.id == administrator.id:
        raise HTTPException(409, "You cannot delete your own account")
    db.execute(delete(UserSession).where(UserSession.user_id == user.id))
    db.delete(user)
    db.commit()


@app.get("/api/v1/doctors")
def doctors(user: User = Depends(require_permission("directory.view")), db: Session = Depends(get_db)):
    stmt = (
        select(Doctor, WaitingRoom)
        .join(WaitingRoom, WaitingRoom.id == Doctor.waiting_room_id)
        .where(Doctor.is_active.is_(True), WaitingRoom.is_active.is_(True))
    )
    if user_access_profile(user, db) == "radiographer":
        if not user.doctor_id:
            return []
        stmt = stmt.where(Doctor.id == user.doctor_id)
    return [
        {
            "id": item.id,
            "name": item.name,
            "department": item.department,
            "designation": item.designation,
            "room": item.room_number,
            "waitingRoom": room.code,
        }
        for item, room in db.execute(stmt.order_by(Doctor.department, Doctor.name)).all()
    ]


@app.get("/api/v1/waiting-rooms")
def waiting_rooms(user: User = Depends(require_permission("directory.view")), db: Session = Depends(get_db)):
    stmt = select(WaitingRoom).where(WaitingRoom.is_active.is_(True))
    if user_access_profile(user, db) == "radiographer":
        doctor = db.get(Doctor, user.doctor_id) if user.doctor_id else None
        if not doctor:
            return []
        stmt = stmt.where(WaitingRoom.id == doctor.waiting_room_id)
    return list(db.scalars(stmt.order_by(WaitingRoom.code)))


@app.get("/api/v1/lookups", response_model=list[LookupRead])
def lookups(
    category: str | None = None,
    include_inactive: bool = False,
    user: User = Depends(require_permission("master_data.view")),
    db: Session = Depends(get_db),
):
    stmt = select(LookupOption)
    if category:
        stmt = stmt.where(LookupOption.category == category)
    if not include_inactive or not has_permission(user, db, "master_data.manage"):
        stmt = stmt.where(LookupOption.is_active.is_(True))
    return list(db.scalars(stmt.order_by(LookupOption.category, LookupOption.sort_order, LookupOption.label)))


@app.post("/api/v1/lookups", response_model=LookupRead, status_code=201)
def create_lookup(payload: LookupCreate, user: User = Depends(require_permission("master_data.manage")), db: Session = Depends(get_db)):
    if db.scalar(
        select(LookupOption).where(LookupOption.category == payload.category, LookupOption.value == payload.value)
    ):
        raise HTTPException(409, "This value already exists in the category")
    item = LookupOption(**payload.model_dump())
    db.add(item)
    db.add(
        AuditEvent(
            action="lookup.created",
            actor=user.username,
            detail={
                "category": payload.category,
                "value": payload.value,
            },
        )
    )
    db.commit()
    db.refresh(item)
    return item


@app.patch("/api/v1/lookups/{lookup_id}", response_model=LookupRead)
def update_lookup(
    lookup_id: str, payload: LookupUpdate, user: User = Depends(require_permission("master_data.manage")), db: Session = Depends(get_db)
):
    item = db.get(LookupOption, lookup_id)
    if not item:
        raise HTTPException(404, "Lookup value not found")
    previous = {"label": item.label, "sort_order": item.sort_order, "is_active": item.is_active}
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    db.add(
        AuditEvent(
            action="lookup.updated",
            actor=user.username,
            detail={
                "category": item.category,
                "value": item.value,
                "previous": previous,
            },
        )
    )
    db.commit()
    db.refresh(item)
    return item


@app.post("/api/v1/admin/waiting-rooms", status_code=201)
def create_waiting_room(
    payload: WaitingRoomCreate, user: User = Depends(require_permission("directory.manage")), db: Session = Depends(get_db)
):
    if db.get(WaitingRoom, payload.id) or db.scalar(select(WaitingRoom).where(WaitingRoom.code == payload.code)):
        raise HTTPException(409, "Waiting room ID or code already exists")
    room = WaitingRoom(**payload.model_dump(), is_active=True)
    db.add(room)
    db.add(AuditEvent(action="waiting_room.created", actor=user.username, detail={"waiting_room": room.code}))
    db.commit()
    db.refresh(room)
    return room


@app.post("/api/v1/admin/doctors", status_code=201)
def create_doctor(payload: DoctorCreate, user: User = Depends(require_permission("directory.manage")), db: Session = Depends(get_db)):
    if db.get(Doctor, payload.id):
        raise HTTPException(409, "Doctor ID already exists")
    if not db.get(WaitingRoom, payload.waiting_room_id):
        raise HTTPException(422, "Waiting room does not exist")
    doctor = Doctor(**payload.model_dump(), is_active=True)
    db.add(doctor)
    db.add(AuditEvent(action="doctor.created", actor=user.username, detail={"doctor_id": doctor.id}))
    db.commit()
    db.refresh(doctor)
    return doctor


@app.patch("/api/v1/admin/doctors/{doctor_id}")
def update_doctor(
    doctor_id: str,
    payload: DoctorUpdate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    doctor = db.get(Doctor, doctor_id)
    if not doctor:
        raise HTTPException(404, "Doctor not found")
    previous = {field: getattr(doctor, field) for field in ("name", "department", "designation", "room_number", "waiting_room_id", "token_prefix")}
    updates = payload.model_dump(exclude_unset=True)
    can_manage_directory = has_permission(user, db, "directory.manage")
    can_manage_room = has_permission(user, db, "doctor.room.manage")
    if not can_manage_directory and not (can_manage_room and set(updates) <= {"room_number"}):
        raise HTTPException(403, "Insufficient permission")
    if not updates:
        return doctor
    if "waiting_room_id" in updates and updates["waiting_room_id"] and not db.get(WaitingRoom, updates["waiting_room_id"]):
        raise HTTPException(422, "Waiting room does not exist")
    for field, value in updates.items():
        if value is not None and field in {"name", "department", "designation", "room_number", "waiting_room_id", "token_prefix"}:
            setattr(doctor, field, value)
    db.add(AuditEvent(action="doctor.updated", actor=user.username, detail={"doctor_id": doctor.id, "previous": previous, "changes": updates}))
    db.commit()
    db.refresh(doctor)
    return doctor


@app.get("/api/v1/holidays", response_model=list[HolidayRead])
def holidays(_: User = Depends(require_permission("holidays.view")), db: Session = Depends(get_db)):
    return list(db.scalars(select(Holiday).where(Holiday.is_active.is_(True)).order_by(Holiday.holiday_date)))


@app.post("/api/v1/holidays", response_model=HolidayRead, status_code=201)
def create_holiday(payload: HolidayCreate, user: User = Depends(require_permission("holidays.manage")), db: Session = Depends(get_db)):
    if db.scalar(select(Holiday).where(Holiday.holiday_date == payload.holiday_date)):
        raise HTTPException(409, "A holiday already exists on this date")
    holiday = Holiday(**payload.model_dump(), is_active=True)
    db.add(holiday)
    db.add(
        AuditEvent(
            action="holiday.created",
            actor=user.username,
            detail={
                "date": payload.holiday_date.isoformat(),
                "name": payload.name,
            },
        )
    )
    db.commit()
    db.refresh(holiday)
    return holiday


@app.patch("/api/v1/holidays/{holiday_id}", response_model=HolidayRead)
def update_holiday(
    holiday_id: str, is_active: bool, user: User = Depends(require_permission("holidays.manage")), db: Session = Depends(get_db)
):
    holiday = db.get(Holiday, holiday_id)
    if not holiday:
        raise HTTPException(404, "Holiday not found")
    holiday.is_active = is_active
    db.add(
        AuditEvent(
            action="holiday.updated",
            actor=user.username,
            detail={
                "date": holiday.holiday_date.isoformat(),
                "is_active": is_active,
            },
        )
    )
    db.commit()
    db.refresh(holiday)
    return holiday


@app.get("/api/v1/patients", response_model=list[PatientRead])
def patients(
    query: str | None = None,
    _: User = Depends(require_permission("patients.view")),
    db: Session = Depends(get_db),
):
    return AppointmentService(db).patients(query)


@app.post("/api/v1/patients", response_model=PatientRead, status_code=201)
def create_patient(
    payload: PatientCreate,
    user: User = Depends(require_permission("patients.manage")),
    db: Session = Depends(get_db),
):
    return AppointmentService(db).create_patient(payload, user.username)


@app.get("/api/v1/schedule-slots", response_model=list[SlotRead])
def schedule_slots(
    doctor_id: str | None = None,
    slot_date: date | None = None,
    user: User = Depends(require_permission("schedule.view")),
    db: Session = Depends(get_db),
):
    if user_access_profile(user, db) == "radiographer":
        doctor_id = user.doctor_id
    return AppointmentService(db).slots(doctor_id, slot_date)


@app.post("/api/v1/schedule-slots", response_model=SlotRead, status_code=201)
def create_schedule_slot(
    payload: SlotCreate,
    user: User = Depends(require_permission("schedule.manage")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, payload.doctor_id, db)
    return AppointmentService(db).create_slot(payload, user.username)


@app.get("/api/v1/appointments", response_model=list[AppointmentRead])
def appointments(
    doctor_id: str | None = None,
    appointment_status: str | None = None,
    user: User = Depends(require_permission("appointments.view")),
    db: Session = Depends(get_db),
):
    if user_access_profile(user, db) == "radiographer":
        doctor_id = user.doctor_id
    return AppointmentService(db).appointments(doctor_id, appointment_status)


@app.post("/api/v1/appointments", response_model=AppointmentRead, status_code=201)
def create_appointment(
    payload: AppointmentCreate,
    user: User = Depends(require_permission("appointments.manage")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, payload.doctor_id, db)
    return AppointmentService(db).create_appointment(payload, user.username)


@app.patch("/api/v1/appointments/{appointment_id}", response_model=AppointmentRead)
def update_appointment(
    appointment_id: str,
    payload: AppointmentUpdate,
    user: User = Depends(require_permission("appointments.manage")),
    db: Session = Depends(get_db),
):
    appointment = db.get(Appointment, appointment_id)
    if not appointment:
        raise HTTPException(404, "Appointment not found")
    enforce_doctor_access(user, appointment.doctor_id, db)
    return AppointmentService(db).update_appointment(appointment_id, payload, user.username)


@app.post("/api/v1/appointments/{appointment_id}/check-in", response_model=TokenRead)
async def check_in_appointment(
    appointment_id: str,
    user: User = Depends(require_permission("appointments.check_in")),
    db: Session = Depends(get_db),
):
    appointment = db.get(Appointment, appointment_id)
    if not appointment:
        raise HTTPException(404, "Appointment not found")
    enforce_doctor_access(user, appointment.doctor_id, db)
    service = AppointmentService(db)
    token = service.check_in(appointment_id, user.username)
    state = QueueService(db).display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room,
        "queue.updated",
        "appointment.checked_in",
        display=state.model_dump(mode="json"),
    )
    return token


@app.get("/api/v1/registration-preview")
def registration_preview(
    user: User = Depends(require_permission("queue.serial.create")), db: Session = Depends(get_db)
):
    from .models import RegistrationCounter
    today = date.today()
    counter = db.get(RegistrationCounter, today.year)
    return {"serial_number": f"{(counter.sequence if counter else 0) + 1:05d}/{today.year % 100:02d}",
            "date": today.isoformat()}


@app.post("/api/v1/tokens", response_model=TokenRead, status_code=201)
async def create_token(
    payload: TokenCreate, user: User = Depends(require_permission("queue.serial.create")), db: Session = Depends(get_db)
):
    service = QueueService(db)
    token = service.create_token(payload, user.username)
    sms = SmsService(db)
    sms.queue(
        token.patient_phone,
        "token_confirmation",
        f"CMH token {token.token_number}. {token.doctor_name}, room {token.room_number}. "
        f"Estimated wait {token.waiting_minutes} minutes.",
    )
    db.commit()
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room, "queue.updated", "token.created", display=state.model_dump(mode="json")
    )
    return token


@app.get("/api/v1/tokens", response_model=list[TokenRead])
def list_tokens(
    doctor_id: str | None = None,
    waiting_room: str | None = None,
    shared_waiting: bool = False,
    user: User = Depends(require_permission("queue.view")),
    db: Session = Depends(get_db),
):
    if user_access_profile(user, db) == "display":
        raise HTTPException(403, "Display accounts cannot access patient lists")
    if user_access_profile(user, db) == "radiographer":
        doctor_id = user.doctor_id
        if not doctor_id:
            return []
        shared_waiting = True
    return QueueService(db).list_tokens(doctor_id, waiting_room, shared_waiting)


@app.post("/api/v1/doctors/{doctor_id}/call-next", response_model=TokenRead)
async def call_next(
    doctor_id: str, user: User = Depends(require_permission("queue.call")), db: Session = Depends(get_db)
):
    enforce_doctor_access(user, doctor_id, db)
    service = QueueService(db)
    token = service.call_next(doctor_id, user.username)
    if has_permission(user, db, "audio.announce"):
        queue_announcement(db, token)
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast_all(
        token.waiting_room, "patient.called", "doctor.call_next", display=state.model_dump(mode="json")
    )
    return token


@app.post("/api/v1/doctors/{doctor_id}/tokens/{token_id}/call", response_model=TokenRead)
async def call_token(
    doctor_id: str,
    token_id: str,
    user: User = Depends(require_permission("queue.call")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, doctor_id, db)
    service = QueueService(db)
    token = service.call(token_id, doctor_id, actor=user.username)
    if has_permission(user, db, "audio.announce"):
        queue_announcement(db, token)
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast_all(
        token.waiting_room, "patient.called", "doctor.manual_call", display=state.model_dump(mode="json")
    )
    return token


@app.post("/api/v1/doctors/{doctor_id}/tokens/{token_id}/complete", response_model=TokenRead)
async def complete_token(
    doctor_id: str,
    token_id: str,
    user: User = Depends(require_permission("queue.action")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, doctor_id, db)
    service = QueueService(db)
    token = service.action(token_id, doctor_id, "complete", None, user.username)
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room, "queue.updated", "consultation.completed", display=state.model_dump(mode="json")
    )
    return token


@app.post("/api/v1/doctors/{doctor_id}/tokens/{token_id}/action", response_model=TokenRead)
async def queue_action(
    doctor_id: str,
    token_id: str,
    payload: QueueActionRequest,
    user: User = Depends(require_permission("queue.action")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, doctor_id, db)
    if payload.action == "recall" and not has_permission(user, db, "audio.announce"):
        raise HTTPException(403, "Audio announcement permission required for recall")
    service = QueueService(db)
    token = service.action(token_id, doctor_id, payload.action, payload.reason, user.username)
    if payload.action == "recall" and has_permission(user, db, "audio.announce"):
        queue_announcement(db, token)
    state = service.display(token.waiting_room)
    event_type = "patient.called" if payload.action == "recall" else "queue.updated"
    broadcast = waiting_room_hub.broadcast_all if payload.action == "recall" else waiting_room_hub.broadcast
    await broadcast(token.waiting_room, event_type, f"queue.{payload.action}", display=state.model_dump(mode="json"))
    return token


@app.patch("/api/v1/tokens/{token_id}/priority", response_model=TokenRead)
async def update_priority(
    token_id: str,
    payload: PriorityUpdateRequest,
    user: User = Depends(require_permission("queue.priority")),
    db: Session = Depends(get_db),
):
    existing = db.get(QueueToken, token_id)
    if not existing:
        raise HTTPException(404, "Token not found")
    if user_access_profile(user, db) == "radiographer" and existing.doctor_id != user.doctor_id:
        raise HTTPException(403, "Radiographer accounts can only manage their assigned queue")
    service = QueueService(db)
    token = service.update_priority(token_id, payload.priority, payload.reason, user.username)
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room, "queue.updated", "queue.priority_changed", display=state.model_dump(mode="json")
    )
    return token


@app.patch("/api/v1/tokens/{token_id}/room", response_model=TokenRead)
async def update_token_room(
    token_id: str,
    payload: TokenRoomUpdateRequest,
    user: User = Depends(require_permission("queue.action")),
    db: Session = Depends(get_db),
):
    existing = db.get(QueueToken, token_id)
    if not existing:
        raise HTTPException(404, "Token not found")
    if user_access_profile(user, db) == "radiographer" and existing.doctor_id != user.doctor_id:
        raise HTTPException(403, "Radiographer accounts can only manage their assigned queue")
    token = QueueService(db).update_room(token_id, payload.room_number, user.username)
    state = QueueService(db).display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room, "queue.updated", "queue.room_changed", display=state.model_dump(mode="json")
    )
    return token


@app.put("/api/v1/doctors/{doctor_id}/queue-control")
async def queue_control(
    doctor_id: str,
    payload: QueueControlUpdate,
    user: User = Depends(require_permission("queue.pause")),
    db: Session = Depends(get_db),
):
    enforce_doctor_access(user, doctor_id, db)
    control = QueueService(db).set_paused(doctor_id, payload.is_paused, payload.reason, user.username)
    doctor = db.get(Doctor, doctor_id)
    room = db.get(WaitingRoom, doctor.waiting_room_id)
    await waiting_room_hub.broadcast(
        room.code, "queue.updated", "queue.paused" if payload.is_paused else "queue.resumed"
    )
    return control


@app.get("/api/v1/doctors/{doctor_id}/queue-control")
def get_queue_control(doctor_id: str, user: User = Depends(require_permission("queue.view")), db: Session = Depends(get_db)):
    enforce_doctor_access(user, doctor_id, db)
    return db.get(QueueControl, doctor_id) or {"doctor_id": doctor_id, "is_paused": False, "reason": None}


@app.post("/api/v1/tokens/{token_id}/transfer", response_model=TokenRead)
async def transfer_token(
    token_id: str,
    payload: TransferRequest,
    user: User = Depends(require_permission("queue.transfer")),
    db: Session = Depends(get_db),
):
    if user_access_profile(user, db) == "radiographer":
        raise HTTPException(403, "Radiographer accounts cannot transfer patients between queues")
    existing = db.get(QueueToken, token_id)
    previous_room = existing.waiting_room if existing else None
    token = QueueService(db).transfer(token_id, payload.doctor_id, payload.reason, user.username)
    service = QueueService(db)
    if previous_room and previous_room != token.waiting_room:
        previous_state = service.display(previous_room)
        await waiting_room_hub.broadcast(
            previous_room,
            "queue.updated",
            "queue.transferred_out",
            display=previous_state.model_dump(mode="json"),
        )
    state = service.display(token.waiting_room)
    await waiting_room_hub.broadcast(
        token.waiting_room, "queue.updated", "queue.transferred", display=state.model_dump(mode="json")
    )
    return token


@app.get("/api/v1/displays/{waiting_room}", response_model=DisplayRead)
def display(waiting_room: str, user: User = Depends(require_permission("display.view")), db: Session = Depends(get_db)):
    if user_access_profile(user, db) == "radiographer":
        raise HTTPException(403, "Radiographer accounts cannot access waiting-room patient displays")
    return QueueService(db).display(waiting_room)


@app.websocket("/api/v1/realtime/waiting-rooms/{waiting_room}")
async def waiting_room_realtime(websocket: WebSocket, waiting_room: str):
    with SessionLocal() as db:
        user = websocket_user(websocket, db)
        if not user:
            await websocket.close(code=4401)
            return
        if not has_permission(user, db, "display.view") or user_access_profile(user, db) == "radiographer":
            await websocket.close(code=4403)
            return
    await waiting_room_hub.connect(waiting_room, websocket)
    try:
        while True:
            message = await websocket.receive_text()
            if message == "ping":
                await websocket.send_json(waiting_room_hub.event(waiting_room, "heartbeat", "pong"))
    except WebSocketDisconnect:
        await waiting_room_hub.disconnect(waiting_room, websocket)


@app.get("/api/v1/dashboard/patients", response_model=list[TokenRead])
def dashboard_patients(
    status: str = "all", user: User = Depends(require_permission("dashboard.view")), db: Session = Depends(get_db)
):
    if status not in {"all", "waiting", "called", "in_progress", "completed", "vip", "wait"}:
        raise HTTPException(422, "Unknown dashboard detail")
    doctor_id = user.doctor_id if user_access_profile(user, db) == "radiographer" else None
    if user_access_profile(user, db) == "radiographer" and not doctor_id:
        return []
    tokens = QueueService(db).list_tokens(doctor_id=doctor_id, shared_waiting=bool(doctor_id))
    if status in {"all", "wait"}:
        return tokens
    if status == "vip":
        return [token for token in tokens if token.priority == "vip"]
    return [token for token in tokens if token.status in ({"called", "recalled"} if status == "called" else {status})]


@app.get("/api/v1/doctors/{doctor_id}/available-patients", response_model=list[TokenRead])
def available_patients(
    doctor_id: str, user: User = Depends(require_permission("queue.action")), db: Session = Depends(get_db)
):
    enforce_doctor_access(user, doctor_id, db)
    return QueueService(db).available_patients(doctor_id)


@app.post("/api/v1/doctors/{doctor_id}/claim/{token_id}", response_model=TokenRead)
async def claim_patient(
    doctor_id: str, token_id: str, user: User = Depends(require_permission("queue.action")), db: Session = Depends(get_db)
):
    enforce_doctor_access(user, doctor_id, db)
    previous = db.get(QueueToken, token_id)
    previous_room = previous.waiting_room if previous else None
    service = QueueService(db)
    token = service.claim_patient(token_id, doctor_id, user.username)
    for room in {previous_room, token.waiting_room} - {None}:
        await waiting_room_hub.broadcast(room, "queue.updated", "queue.reassigned", display=service.display(room).model_dump(mode="json"))
    return token


@app.get("/api/v1/dashboard", response_model=DashboardRead)
def dashboard(user: User = Depends(require_permission("dashboard.view")), db: Session = Depends(get_db)):
    today = date.today()
    assigned_only = user_access_profile(user, db) == "radiographer"
    doctor_stmt = (
        select(Doctor, WaitingRoom)
        .join(WaitingRoom, WaitingRoom.id == Doctor.waiting_room_id)
        .where(Doctor.is_active.is_(True), WaitingRoom.is_active.is_(True))
    )
    if assigned_only:
        doctor_stmt = doctor_stmt.where(Doctor.id == user.doctor_id)
    doctors = list(db.execute(doctor_stmt.order_by(Doctor.department, Doctor.name)).all())
    doctor_ids = [doctor.id for doctor, _ in doctors]
    token_stmt = select(QueueToken).where(QueueToken.token_date == today)
    if assigned_only:
        token_stmt = token_stmt.where(or_(QueueToken.doctor_id == user.doctor_id, QueueToken.status == "waiting"))
    tokens = list(db.scalars(token_stmt)) if doctor_ids or not assigned_only else []

    def metrics(items: list[QueueToken]) -> tuple[dict[str, int], int]:
        counts: dict[str, int] = {}
        waits = []
        for token in items:
            counts[token.status] = counts.get(token.status, 0) + 1
            if token.called_at:
                waits.append((token.called_at - token.created_at).total_seconds())
        average = round(sum(waits) / len(waits) / 60) if waits else 0
        return counts, average

    counts, average_wait = metrics(tokens)
    radiographers = []
    for doctor, room in doctors:
        doctor_tokens = [token for token in tokens if token.doctor_id == doctor.id]
        doctor_counts, doctor_wait = metrics(doctor_tokens)
        radiographers.append(DashboardDoctorRead(
            doctor_id=doctor.id, doctor_name=doctor.name, department=doctor.department,
            room_number=doctor.room_number, waiting_room=room.code,
            waiting=doctor_counts.get("waiting", 0),
            called=doctor_counts.get("called", 0) + doctor_counts.get("recalled", 0),
            in_progress=doctor_counts.get("in_progress", 0), completed=doctor_counts.get("completed", 0),
            total=len(doctor_tokens), average_wait_minutes=doctor_wait,
            vip=sum(token.priority == "vip" for token in doctor_tokens),
        ))
    vip_tokens = [token for token in tokens if token.priority == "vip"]

    all_room_numbers = set()
    for doctor, _ in doctors:
        if doctor.room_number:
            all_room_numbers.add(doctor.room_number)
    for token in tokens:
        if token.room_number:
            all_room_numbers.add(token.room_number)
    if not assigned_only:
        lookup_rooms = db.scalars(
            select(LookupOption.value).where(LookupOption.category == "room_number", LookupOption.is_active.is_(True))
        ).all()
        for lr in lookup_rooms:
            if lr:
                all_room_numbers.add(lr)

    def room_sort_key(val: str):
        digits = "".join(ch for ch in val if ch.isdigit())
        return (int(digits) if digits else 9999, val)

    sorted_rooms = sorted(all_room_numbers, key=room_sort_key)
    room_stats = []
    for r_num in sorted_rooms:
        r_doctors = [doc for doc, _ in doctors if doc.room_number == r_num]
        r_doc_ids = {doc.id for doc in r_doctors}
        r_tokens = [
            token for token in tokens
            if token.room_number == r_num or (not token.room_number and token.doctor_id in r_doc_ids)
        ]
        r_counts, r_wait = metrics(r_tokens)
        r_doc_name = ", ".join(d.name for d in r_doctors) if r_doctors else ""
        r_dept = ", ".join(sorted({d.department for d in r_doctors if d.department})) if r_doctors else ""
        r_wr = ", ".join(sorted({rm.code for doc, rm in doctors if doc.room_number == r_num})) if r_doctors else ""

        room_stats.append(DashboardRoomRead(
            room_number=r_num,
            doctor_name=r_doc_name,
            department=r_dept,
            waiting_room=r_wr,
            waiting=r_counts.get("waiting", 0),
            called=r_counts.get("called", 0) + r_counts.get("recalled", 0),
            in_progress=r_counts.get("in_progress", 0),
            completed=r_counts.get("completed", 0),
            total=len(r_tokens),
            average_wait_minutes=r_wait,
            vip=sum(token.priority == "vip" for token in r_tokens),
        ))

    return DashboardRead(
        scope="assigned" if assigned_only else "all", generated_at=datetime.utcnow(), total=len(tokens),
        waiting=counts.get("waiting", 0),
        called=counts.get("called", 0) + counts.get("recalled", 0),
        in_progress=counts.get("in_progress", 0),
        completed=counts.get("completed", 0),
        average_wait_minutes=average_wait,
        vip_total=len(vip_tokens),
        vip_waiting=sum(token.status == "waiting" for token in vip_tokens),
        vip_active=sum(token.status in {"called", "recalled", "in_progress"} for token in vip_tokens),
        vip_completed=sum(token.status == "completed" for token in vip_tokens),
        radiographers=radiographers,
        rooms=room_stats,
    )


@app.get("/api/v1/settings")
def settings(_: User = Depends(require_permission("settings.manage")), db: Session = Depends(get_db)):
    return list(db.scalars(select(AppSetting).order_by(AppSetting.key)))


@app.put("/api/v1/settings/{key}")
def update_setting(
    key: str, payload: SettingUpdate, _: User = Depends(require_permission("settings.manage")), db: Session = Depends(get_db)
):
    setting = db.get(AppSetting, key)
    if not setting:
        raise HTTPException(404, "Setting not found")
    setting.value = validated_setting(key, {**setting.value, **payload.value})
    db.commit()
    db.refresh(setting)
    return setting


@app.post("/api/v1/settings/announcement/test", status_code=202)
def test_announcement_voice(
    payload: SettingUpdate, _: User = Depends(require_permission("settings.manage"))
):
    values = validated_setting("announcement", payload.value)
    sample = type(
        "AudioPreview",
        (),
        {
            "token_number": "CMH-123",
            "patient_name": "Rahim Uddin",
            "doctor_name": "Dr. Ayesha Khan",
            "room_number": "205",
        },
    )()
    announcement_engine.enqueue(sample, values)
    return {"status": "queued", "voice_mode": values.get("voice_mode", "auto")}


@app.get("/api/v1/audit", response_model=list[AuditRead])
def audit(limit: int = 100, _: User = Depends(require_permission("reports.view")), db: Session = Depends(get_db)):
    return list(db.scalars(select(AuditEvent).order_by(AuditEvent.created_at.desc()).limit(min(max(limit, 1), 500))))


@app.get("/api/v1/reports/queue-summary")
def queue_report(_: User = Depends(require_permission("reports.view")), db: Session = Depends(get_db)):
    tokens = list(db.scalars(select(QueueToken).where(QueueToken.token_date == date.today())))
    def grouped(field: str) -> dict[str, int]:
        result: dict[str, int] = {}
        for token in tokens:
            key = str(getattr(token, field) or "Unassigned")
            result[key] = result.get(key, 0) + 1
        return result

    wait_seconds = [(token.called_at - token.created_at).total_seconds() for token in tokens if token.called_at]
    service_seconds = [
        (token.completed_at - token.started_at).total_seconds()
        for token in tokens
        if token.completed_at and token.started_at
    ]
    return {
        "total": len(tokens),
        "by_status": grouped("status"),
        "by_doctor": grouped("doctor_name"),
        "by_department": grouped("department"),
        "by_room": grouped("waiting_room"),
        "by_source": grouped("source"),
        "by_priority": grouped("priority"),
        "priority": sum(token.priority != "normal" for token in tokens),
        "average_wait_minutes": round(sum(wait_seconds) / len(wait_seconds) / 60, 1) if wait_seconds else 0,
        "longest_wait_minutes": round(max(wait_seconds) / 60, 1) if wait_seconds else 0,
        "average_service_minutes": round(sum(service_seconds) / len(service_seconds) / 60, 1) if service_seconds else 0,
    }


def reception_report_query(
    db: Session,
    date_from: date,
    date_to: date,
    doctor_id: str | None,
    waiting_room: str | None,
    status_filter: str | None,
    priority: str | None,
    service_category: str | None,
) -> list[QueueToken]:
    if date_from > date_to:
        raise HTTPException(422, "From date must not be after to date")
    if (date_to - date_from).days > 366:
        raise HTTPException(422, "Report range cannot exceed 366 days")
    statement = select(QueueToken).where(QueueToken.token_date.between(date_from, date_to))
    for field, value in (
        (QueueToken.doctor_id, doctor_id),
        (QueueToken.waiting_room, waiting_room),
        (QueueToken.status, status_filter),
        (QueueToken.priority, priority),
        (QueueToken.service_category, service_category),
    ):
        if value:
            statement = statement.where(field == value)
    return list(db.scalars(statement.order_by(QueueToken.token_date.desc(), QueueToken.created_at.desc())))


def reception_report_params(
    date_from: date = Query(default_factory=date.today),
    date_to: date = Query(default_factory=date.today),
    doctor_id: str | None = None,
    waiting_room: str | None = None,
    status: str | None = None,
    priority: str | None = None,
    service_category: str | None = None,
) -> dict[str, object]:
    return {
        "date_from": date_from, "date_to": date_to, "doctor_id": doctor_id,
        "waiting_room": waiting_room, "status_filter": status,
        "priority": priority, "service_category": service_category,
    }


@app.get("/api/v1/reports/reception")
def reception_report(
    filters: dict[str, object] = Depends(reception_report_params),
    _: User = Depends(require_permission("reports.view")),
    db: Session = Depends(get_db),
):
    return reception_report_query(db, **filters)


@app.get("/api/v1/reports/reception.xlsx")
def reception_report_excel(
    filters: dict[str, object] = Depends(reception_report_params),
    _: User = Depends(require_permission("reports.view")),
    db: Session = Depends(get_db),
):
    tokens = reception_report_query(db, **filters)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reception Desk"
    labels = {(item.category, item.value): item.label for item in db.scalars(select(LookupOption))}
    def label(category, value):
        return labels.get((category, value), value or "")
    headers = [
        "Serial", "Date", "Service No./BA", "Designation / Rank", "Name", "Age", "Unit",
        "MRI Area", "Contrast", "Film", "Report", "Patient Source", "Priority", "Room", "Radiographer", "Status",
        "Created At", "Called At", "Service Started At", "Completed At", "Cancelled At", "Skipped At",
        "Recalled At", "No Show At", "Scheduled At", "Recall Count", "Mobile Number", "Service Category",
        "Department", "Waiting Room", "Registration Type", "Patient Type", "Entitlement", "Service Status", "Sponsor Rank", "Family Relationship", "Summary Category",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="245B45")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for token in tokens:
        sheet.append([
            token.serial_number or token.token_number, token.token_date, token.service_number or "",
            label("rank_relationship", token.rank), token.patient_name, token.age, token.unit or "",
            token.mri_area or "", token.contrast, token.film, token.report or "", label("patient_source", token.patient_source),
            label("priority_category", token.priority), token.room_number, token.doctor_name, token.status,
            token.created_at, token.called_at, token.started_at, token.completed_at, token.cancelled_at,
            token.skipped_at, token.recalled_at, token.no_show_at, token.scheduled_at, token.recall_count,
            token.patient_phone, token.service_category, token.department, token.waiting_room, token.source,
            label("beneficiary_type", token.beneficiary_type), label("entitlement", token.entitlement), label("service_status", token.service_status),
            label("rank_relationship", token.sponsor_rank), label("family_relationship", token.family_relationship), token.summary_category or "Needs review",
        ])
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [16, 13, 20, 25, 28, 8, 22, 28, 10, 8, 45, 22, 16, 12, 28, 16] + [20] * 9 + [12, 20, 20, 20, 16, 18] + [22] * 6
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        for cell in row[16:25]:
            cell.number_format = "yyyy-mm-dd hh:mm"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Patient-entered text must never execute as an Excel formula.
    output = BytesIO()
    workbook.save(output)
    filename = f"reception-report-{filters['date_from']}-to-{filters['date_to']}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/v1/reports/reception.pdf")
def reception_report_pdf(
    filters: dict[str, object] = Depends(reception_report_params),
    _: User = Depends(require_permission("reports.view")),
    db: Session = Depends(get_db),
):
    tokens = reception_report_query(db, **filters)
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=8 * mm, rightMargin=8 * mm,
        topMargin=9 * mm, bottomMargin=9 * mm,
        title="Reception Desk Report", author="CMH Smart Serial",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Reception Desk Report", styles["Title"]),
        Paragraph(
            f"Date range: {filters['date_from']} to {filters['date_to']} &nbsp;&nbsp; Records: {len(tokens)}",
            styles["Normal"],
        ),
        Spacer(1, 4 * mm),
    ]
    labels = {(item.category, item.value): item.label for item in db.scalars(select(LookupOption))}
    def label(category, value):
        return labels.get((category, value), value or "")
    headers = ["Serial", "Date", "Service No./BA", "Designation", "Name", "Age", "Unit", "MRI area", "Contrast", "Film", "Report", "Patient source", "Priority / Room"]
    rows = [headers]
    cell_style = styles["BodyText"].clone("RegisterCell")
    cell_style.fontSize = 6.5
    cell_style.leading = 8
    for token in tokens:
        values = [token.serial_number or token.token_number, str(token.token_date), token.service_number or "",
            label("rank_relationship", token.rank), token.patient_name, token.age, token.unit or "", token.mri_area or "",
            token.contrast, token.film, token.report or "", label("patient_source", token.patient_source),
            f"{label('priority_category', token.priority)} / {token.room_number}"]
        rows.append([Paragraph(escape(str(value)) if value is not None else "", cell_style) for value in values])
    table = LongTable(rows, repeatRows=1, splitInRow=1,
        colWidths=[18*mm, 17*mm, 21*mm, 23*mm, 28*mm, 9*mm, 19*mm, 25*mm, 12*mm, 9*mm, 44*mm, 21*mm, 23*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123C31")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("LEADING", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#B8C5C0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6F4")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    document.build(story)
    filename = f"reception-report-{filters['date_from']}-to-{filters['date_to']}.pdf"
    return Response(
        content=output.getvalue(), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/v1/devices", response_model=list[DeviceRead])
def devices(_: User = Depends(require_permission("devices.view")), db: Session = Depends(get_db)):
    return list(db.scalars(select(DeviceEndpoint).order_by(DeviceEndpoint.waiting_room, DeviceEndpoint.name)))


@app.post("/api/v1/devices", status_code=201)
def create_device(payload: DeviceCreate, user: User = Depends(require_permission("devices.manage")), db: Session = Depends(get_db)):
    if not db.scalar(
        select(WaitingRoom).where(WaitingRoom.code == payload.waiting_room, WaitingRoom.is_active.is_(True))
    ):
        raise HTTPException(422, "Waiting room does not exist")
    raw_key = secrets.token_urlsafe(32)
    device = DeviceEndpoint(**payload.model_dump(), client_key_hash=hashlib.sha256(raw_key.encode()).hexdigest())
    db.add(device)
    db.add(
        AuditEvent(
            action="device.created",
            actor=user.username,
            detail={"device_id": device.id, "type": payload.device_type, "waiting_room": payload.waiting_room},
        )
    )
    db.commit()
    db.refresh(device)
    return {"device": DeviceRead.model_validate(device), "client_key": raw_key}


def authenticated_device(device_id: str, client_key: str | None, db: Session) -> DeviceEndpoint:
    device = db.get(DeviceEndpoint, device_id)
    digest = hashlib.sha256((client_key or "").encode()).hexdigest()
    if not device or not device.is_active or not secrets.compare_digest(device.client_key_hash, digest):
        raise HTTPException(401, "Invalid device credentials")
    return device


@app.post("/api/v1/devices/{device_id}/heartbeat", response_model=DeviceRead)
def device_heartbeat(
    device_id: str,
    fault: str | None = None,
    x_device_key: str | None = Header(default=None),
    db: Session = Depends(get_db),
):
    device = authenticated_device(device_id, x_device_key, db)
    device.last_heartbeat_at = datetime.utcnow()
    device.fault = fault
    db.commit()
    db.refresh(device)
    return device


@app.post("/api/v1/devices/{device_id}/acknowledge", response_model=DeviceRead)
def device_acknowledge(
    device_id: str, event_id: str, x_device_key: str | None = Header(default=None), db: Session = Depends(get_db)
):
    device = authenticated_device(device_id, x_device_key, db)
    device.last_heartbeat_at = datetime.utcnow()
    device.last_acknowledged_event_id = event_id
    device.fault = None
    db.commit()
    db.refresh(device)
    return device


@app.get("/api/v1/notifications/sms")
def sms_messages(limit: int = 100, _: User = Depends(require_permission("sms.view")), db: Session = Depends(get_db)):
    return list(db.scalars(select(SmsMessage).order_by(SmsMessage.created_at.desc()).limit(min(max(limit, 1), 500))))


@app.post("/api/v1/notifications/sms/{message_id}/retry")
def retry_sms(message_id: str, _: User = Depends(require_permission("sms.retry")), db: Session = Depends(get_db)):
    message = db.get(SmsMessage, message_id)
    if not message:
        raise HTTPException(404, "SMS message not found")
    message.status = "queued"
    message.error = None
    message.attempts = 0
    message.next_attempt_at = None
    message.claimed_at = None
    db.commit()
    return message


# A production Angular build is served by the same FastAPI process and port.
# API and WebSocket routes are registered first, so this mount is only the UI.
app.include_router(monthly_report_router)

frontend_dist = Path(
    os.getenv(
        "CMH_SMS_FRONTEND_DIST",
        str(Path(__file__).resolve().parents[2] / "frontend" / "dist" / "cmh-smart-serial" / "browser"),
    )
)
if frontend_dist.is_dir():
    app.mount("/", StaticFiles(directory=frontend_dist, html=True), name="frontend")
