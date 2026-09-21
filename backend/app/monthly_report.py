"""MRI patient-state report. Categories are snapshots, never inferred for legacy rows."""
from __future__ import annotations

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from .auth import require_permission
from .database import get_db
from pydantic import BaseModel, ConfigDict

from .models import AppSetting, AuditEvent, LookupOption, QueueToken, User
from .schemas import LookupRead, ReportClassification, TokenRead

COLUMNS = [
    ('serving_officer', 'SERVING', "OFFR’S + AFNS"),
    ('serving_cadet', 'SERVING', 'OFF / NURSING CADET'),
    ('serving_jco_or', 'SERVING', 'JCO’S / OR / RECT'),
    ('serving_nce', 'SERVING', 'NCE'),
    ('serving_civil', 'SERVING', 'ALL CIV ENTD'),
    ('retired_officer', 'RETD (MIL)', 'OFFR’S / OFFR’S FAMILY'),
    ('retired_or', 'RETD (MIL)', 'OR’S'),
    ('family_officer', 'FAMILY', 'OFFR’S'),
    ('family_jco_or_nce', 'FAMILY', 'JCO’S / OR’S & NCE'),
    ('family_civil', 'FAMILY', 'ALL CIV ENTD'),
    ('re', '', 'RE'), ('cne', '', 'CNE'),
]
GROUPS = [('officer', 'Officers / AFNS'), ('cadet', 'Officer / Nursing cadet'),
          ('jco', 'JCO'), ('or', 'OR / Recruit'), ('nce', 'NCE')]
CLASSIFICATION_LOOKUPS = {
    'beneficiary_type': [('self', 'Self'), ('family', 'Family')],
    'service_status': [('serving', 'Serving'), ('retired', 'Retired')],
    'entitlement': [('military', 'Military'), ('civil', 'Civil entitled'), ('re', 'RE'), ('cne', 'CNE')],
    'family_relationship': [('spouse', 'Spouse'), ('child', 'Child'), ('parent', 'Parent'), ('other', 'Other dependent')],
}
RANK_GROUPS = dict.fromkeys(['brigadier_general', 'colonel', 'lieutenant_colonel', 'major', 'captain', 'lieutenant', 'officer'], 'officer')
RANK_GROUPS.update(afns='officer', cadet='cadet', nce='nce', warrant_officer='jco', jco='jco', sergeant='or', corporal='or', shoinik='or', soldier='or')
router = APIRouter(prefix='/api/v1', tags=['Monthly MRI summary'])


MAPPING_KEY = 'mri_summary_mapping'


def mapping_rows(db):
    options = list(db.scalars(select(LookupOption).where(
        LookupOption.category.in_([*CLASSIFICATION_LOOKUPS, 'rank_relationship']),
        LookupOption.is_active.is_(True),
    ).order_by(LookupOption.sort_order, LookupOption.label)))

    def labels(category):
        grouped = {}
        for option in options:
            if option.category == category:
                code = (option.metadata_json or {}).get('report_code', option.value)
                grouped.setdefault(code, []).append(option.label)
        return {code: ' / '.join(names) for code, names in grouped.items()}

    people, statuses, entitlements = (labels(category) for category in
                                      ('beneficiary_type', 'service_status', 'entitlement'))
    rows = [{'key': key, 'label': f'Entitlement: {entitlements[key]}', 'default_category': key}
            for key in ('re', 'cne') if key in entitlements]
    for person in ('self', 'family'):
        if person not in people:
            continue
        if 'civil' in entitlements:
            rows.append({'key': f'civil:{person}', 'label': f"{entitlements['civil']} · {people[person]}",
                         'default_category': 'family_civil' if person == 'family' else 'serving_civil'})
        if 'military' not in entitlements:
            continue
        for status in ('serving', 'retired'):
            if status not in statuses:
                continue
            prefix = f"{entitlements['military']} · {people[person]} · {statuses[status]}"
            suffix = ' (sponsor)' if person == 'family' else ''
            for group, label in GROUPS:
                rows.append({'key': f'military:{person}:{status}:{group}',
                             'label': f'{prefix} · {label}{suffix}',
                             'default_category': default_military_category(person, status, group)})
            for rank in options:
                if rank.category == 'rank_relationship':
                    rows.append({'key': f'rank:{person}:{status}:{rank.value}',
                                 'label': f'{prefix} · {rank.label}{suffix}',
                                 'default_category': '__inherit__'})
    return rows


def configured_category(db, key, default):
    setting = db.get(AppSetting, MAPPING_KEY)
    category = setting.value.get(key, default) if setting else default
    return default if category == '__inherit__' else category


CATEGORY_LABELS = {
    'serving_officer': ('Serving · Officers / AFNS', 'Serving military officers and AFNS.'),
    'serving_cadet': ('Serving · Officer / Nursing cadets', 'Serving officer cadets and nursing cadets.'),
    'serving_jco_or': ('Serving · JCOs / Other ranks / Recruits', 'Serving JCOs, other ranks and recruits.'),
    'serving_nce': ('Serving · NCE', 'Serving NCE patients.'),
    'serving_civil': ('Serving · Civil entitled', 'Civil-entitled patients classified as Self.'),
    'retired_officer': ('Retired military · Officers / Officer family', 'Retired officers and their eligible family classifications.'),
    'retired_or': ('Retired military · Other ranks', 'Retired other-rank patients.'),
    'family_officer': ('Family · Officers', 'Family patients in the officer category.'),
    'family_jco_or_nce': ('Family · JCOs / Other ranks / NCE', 'Family patients in the JCO, other-rank or NCE category.'),
    'family_civil': ('Family · Civil entitled', 'Civil-entitled family patients.'),
    're': ('RE', 'The RE column on the approved MRI summary form.'),
    'cne': ('CNE', 'The CNE column on the approved MRI summary form.'),
}


def category_options():
    return [{'key': key, 'label': CATEGORY_LABELS[key][0], 'description': CATEGORY_LABELS[key][1],
             'report_label': f'{group} · {label}' if group else label} for key, group, label in COLUMNS]


@router.get('/classification-options')
def classification_options(db: Session = Depends(get_db), _user: User = Depends(require_permission('queue.serial.create'))):
    from .registration import requirements
    lookups = db.scalars(select(LookupOption).where(
        LookupOption.category.in_([*CLASSIFICATION_LOOKUPS, 'rank_relationship'])
    ).order_by(LookupOption.sort_order, LookupOption.label))
    return {'columns': category_options(), 'registration': requirements(db),
            'lookups': [LookupRead.model_validate(item).model_dump() for item in lookups]}


def mapping_response(db):
    setting = db.get(AppSetting, MAPPING_KEY)
    saved = setting.value if setting else {}
    return {'rows': [{**row, 'category': saved.get(row['key'], row['default_category'])} for row in mapping_rows(db)],
            'columns': category_options()}


class MappingUpdate(BaseModel):
    mappings: dict[str, str | None]


@router.get('/mri-summary-mapping')
def get_mapping(db: Session = Depends(get_db), _user: User = Depends(require_permission('settings.manage'))):
    return mapping_response(db)


@router.put('/mri-summary-mapping')
def save_mapping(payload: MappingUpdate, db: Session = Depends(get_db), user: User = Depends(require_permission('settings.manage'))):
    valid_keys = {row['key'] for row in mapping_rows(db)}
    valid_categories = {key for key, _, _ in COLUMNS} | {None}
    if set(payload.mappings) != valid_keys:
        raise HTTPException(409, 'Dropdown options changed. Review the refreshed mappings and save again.')
    if any(value not in valid_categories and not (key.startswith('rank:') and value == '__inherit__') for key, value in payload.mappings.items()):
        raise HTTPException(422, 'Provide every mapping with a valid report category or Needs review')
    setting = db.get(AppSetting, MAPPING_KEY)
    previous = setting.value if setting else {}
    if not setting:
        setting = AppSetting(key=MAPPING_KEY, description='MRI summary classification mapping', value={})
        db.add(setting)
    setting.value = {**previous, **payload.mappings}
    db.add(AuditEvent(action='mri_summary.mapping_updated', actor=user.username,
                      detail={'previous': previous, 'changes': payload.mappings}))
    db.commit()
    return mapping_response(db)


def classify(db: Session, payload) -> str | None:
    """Resolve only documented combinations. Unknown combinations stay visible for review."""
    values = {}
    for field in CLASSIFICATION_LOOKUPS:
        value = getattr(payload, field, None)
        if not value:
            values[field] = None
            continue
        option = db.scalar(select(LookupOption).where(LookupOption.category == field, LookupOption.value == value, LookupOption.is_active.is_(True)))
        if not option:
            raise HTTPException(422, f'Unknown or inactive {field.replace("_", " ")}')
        values[field] = (option.metadata_json or {}).get('report_code', value)
    if values['entitlement'] in ('re', 'cne'):
        return configured_category(db, values['entitlement'], values['entitlement'])
    person = values['beneficiary_type']
    if person not in ('self', 'family'):
        return None
    if person == 'family' and not values['family_relationship']:
        return None
    if values['entitlement'] == 'civil':
        return configured_category(db, f'civil:{person}', 'family_civil' if person == 'family' else 'serving_civil')
    if values['entitlement'] != 'military':
        return None
    rank = payload.sponsor_rank if person == 'family' else getattr(payload, 'rank', None)
    option = db.scalar(select(LookupOption).where(LookupOption.category == 'rank_relationship', LookupOption.value == rank, LookupOption.is_active.is_(True))) if rank else None
    if person == 'family' and not option:
        if rank:
            raise HTTPException(422, 'Select an active sponsor rank for a military family patient')
        return None
    group = (option.metadata_json or {}).get('report_group') if option else None
    status = values['service_status']
    key = f'military:{person}:{status}:{group}'
    fallback = configured_category(db, key, default_military_category(person, status, group))
    return configured_category(db, f'rank:{person}:{status}:{rank}', fallback) if option and status in ('serving', 'retired') else fallback


def default_military_category(person, status, group):
    if status == 'retired':
        if group == 'officer':
            return 'retired_officer'
        if person == 'self' and group == 'or':
            return 'retired_or'
        return None  # The supplied sheet does not establish other retired-family mappings.
    if status != 'serving':
        return None
    if person == 'family':
        return 'family_officer' if group == 'officer' else 'family_jco_or_nce' if group in ('jco', 'or', 'nce') else None
    return {'officer': 'serving_officer', 'cadet': 'serving_cadet', 'jco': 'serving_jco_or', 'or': 'serving_jco_or', 'nce': 'serving_nce'}.get(group)


def month_rows(db, month, basis, date_from=None, date_to=None):
    if date_from is not None or date_to is not None:
        if not date_from or not date_to or date_from > date_to or date_to == date.max:
            raise HTTPException(422, "Provide a valid from/to date range")
        if (date_to - date_from).days > 366:
            raise HTTPException(422, "Select a range of at most 367 days")
        start, end = date_from, date_to + timedelta(days=1)
    else:
        start, end = month_bounds(month)
    return range_rows(db, start, end, basis)


def month_bounds(month):
    try:
        start = date.fromisoformat((month or '') + '-01')
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    except (ValueError, OverflowError):
        raise HTTPException(422, 'Month must use YYYY-MM and precede December 9999')
    return start, end


def range_rows(db, start, end, basis):
    stmt = select(QueueToken)
    if basis == 'completed':
        zone = ZoneInfo('Asia/Dhaka')
        def utc_boundary(day):
            return datetime.combine(day, time.min, zone).astimezone(timezone.utc).replace(tzinfo=None)
        stmt = stmt.where(QueueToken.status == 'completed', QueueToken.completed_at >= utc_boundary(start), QueueToken.completed_at < utc_boundary(end))
    else:
        stmt = stmt.where(QueueToken.token_date >= start, QueueToken.token_date < end)
    records = list(db.scalars(stmt.order_by(QueueToken.token_date, QueueToken.sequence)))
    return start, end, records


def record_day(token, basis):
    if basis == 'completed':
        return token.completed_at.replace(tzinfo=timezone.utc).astimezone(ZoneInfo('Asia/Dhaka')).date()
    return token.token_date


def summary(db, month, basis, date_from=None, date_to=None):
    start, end, records = month_rows(db, month, basis, date_from, date_to)
    keys = [c[0] for c in COLUMNS] + ['unclassified']
    days = { (start + timedelta(days=i)).isoformat(): dict.fromkeys(keys, 0) for i in range((end-start).days)}
    for token in records:
        key = token.summary_category if token.summary_category in keys else 'unclassified'
        days[record_day(token, basis).isoformat()][key] += 1
    rows = [{'date': day, 'counts': counts, 'total': sum(counts.values())} for day, counts in days.items()]
    totals = {key: sum(row['counts'][key] for row in rows) for key in keys}
    return {'date_from': start.isoformat(), 'date_to': (end - timedelta(days=1)).isoformat(), 'month': month, 'basis': basis, 'columns': [{'key': key, 'group': group, 'label': label} for key, group, label in COLUMNS], 'rows': rows, 'totals': totals, 'total': len(records)}


@router.get('/reports/mri-summary')
def get_summary(month: str | None = Query(default=None, pattern=r'^\d{4}-\d{2}$'), date_from: date | None = None, date_to: date | None = None, basis: Literal['registrations', 'completed'] = 'registrations', db: Session = Depends(get_db), _user: User = Depends(require_permission('reports.view'))):
    return summary(db, month, basis, date_from, date_to)


@router.get('/reports/mri-summary/patients', response_model=list[TokenRead])
def summary_patients(month: str | None = Query(default=None, pattern=r'^\d{4}-\d{2}$'), date_from: date | None = None, date_to: date | None = None, basis: Literal['registrations', 'completed'] = 'registrations', day: date | None = None, category: str | None = None, db: Session = Depends(get_db), _user: User = Depends(require_permission('reports.view'))):
    _, _, records = month_rows(db, month, basis, date_from, date_to)
    keys = {c[0] for c in COLUMNS}
    return [token for token in records if (not day or record_day(token, basis) == day) and (not category or (token.summary_category if token.summary_category in keys else 'unclassified') == category)]


class ClassificationCorrection(ReportClassification):
    model_config = ConfigDict(extra='forbid')
    rank: str | None = None
    mode: Literal['inputs', 'direct'] = 'inputs'
    summary_category: str | None = None


@router.patch('/tokens/{token_id}/classification', response_model=TokenRead)
def correct_classification(token_id: str, payload: ClassificationCorrection, db: Session = Depends(get_db), user: User = Depends(require_permission('queue.serial.create'))):
    token = db.scalar(select(QueueToken).where(QueueToken.id == token_id).with_for_update())
    if not token:
        raise HTTPException(404, 'Patient not found')
    input_keys = {*ReportClassification.model_fields, 'rank'}
    before = {key: getattr(token, key) for key in input_keys}
    before.update(summary_category=token.summary_category, summary_category_source=token.summary_category_source)
    if payload.mode == 'direct':
        if 'summary_category' not in payload.model_fields_set or payload.summary_category not in {key for key, _, _ in COLUMNS} | {None}:
            raise HTTPException(422, 'Choose an MRI summary category or Needs review')
        if payload.model_fields_set & input_keys:
            raise HTTPException(422, 'Direct classification accepts only the report category; use patient inputs to edit details')
        token.summary_category = payload.summary_category
        token.summary_category_source = 'manual'
    else:
        if 'summary_category' in payload.model_fields_set:
            raise HTTPException(422, 'Use direct classification to select a report category')
        from .registration import validate_registration
        supplied = {key: value for key, value in payload.model_dump(exclude_unset=True).items() if key in input_keys}
        # Validate only submitted input, preserving omitted/disabled historical values.
        validated = validate_registration(db, ClassificationCorrection(**supplied), existing=token, check_required=False)
        changes = {key: getattr(validated, key) for key in input_keys if key in validated.model_fields_set}
        merged = ClassificationCorrection(**{key: changes.get(key, getattr(token, key)) for key in input_keys})
        # Required dependent fields must be checked against the complete classification.
        from .registration import applicable, blank, requirements, FIELDS
        policy = requirements(db)
        missing = [FIELDS[key] for key in policy['required'] if key in input_keys and key in policy['enabled']
                   and applicable(db, merged, key) and blank(getattr(merged, key))]
        if missing:
            raise HTTPException(422, 'Required fields: ' + ', '.join(missing))
        token.summary_category = classify(db, merged)
        token.summary_category_source = 'automatic'
        for key, value in changes.items():
            setattr(token, key, value)
    after = {key: getattr(token, key) for key in before}
    db.add(AuditEvent(action='token.classification_updated', actor=user.username, token_id=token.id,
                      detail={'mode': payload.mode, 'before': before, 'after': after}))
    db.commit()
    return token


@router.get('/reports/mri-summary.{format}')
def export_summary(format: Literal['xlsx', 'pdf'], month: str | None = Query(default=None, pattern=r'^\d{4}-\d{2}$'), date_from: date | None = None, date_to: date | None = None, basis: Literal['registrations', 'completed'] = 'registrations', db: Session = Depends(get_db), _user: User = Depends(require_permission('reports.view'))):
    report = summary(db, month, basis, date_from, date_to)
    keys = [c[0] for c in COLUMNS]
    review = report['totals']['unclassified'] > 0
    if review:
        keys += ['unclassified']
    period = f"{report['date_from']} to {report['date_to']}"
    title = f'PATIENT STATE {period}'
    note = ('Registrations (including cancelled registrations)' if basis == 'registrations' else 'Completed MRIs by completion date (Asia/Dhaka)')
    if review:
        note += ' | Needs review: included in total, pending classification'
    first = ['DATE', 'SERVING', '', '', '', '', 'RETD (MIL)', '', 'FAMILY', '', '', 'RE', 'CNE'] + (['NEEDS REVIEW'] if review else []) + ['TOTAL']
    second = [''] + [c[2] for c in COLUMNS[:10]] + ['', ''] + ([''] if review else []) + ['']
    data = [first, second]
    for row in report['rows']:
        data.append([date.fromisoformat(row['date']).strftime('%d.%m.%y')] + [row['counts'][key] for key in keys] + [row['total']])
    data.append(['TOTAL'] + [report['totals'][key] for key in keys] + [report['total']])
    spans = [(1, 5), (6, 7), (8, 10)]
    out = BytesIO()
    if format == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'MRI monthly summary'
        for text in ['DEPARTMENT OF RADIOLOGY & IMAGING', 'CMH DHAKA', title, 'MRI CENTRE', note]:
            sheet.append([text])
            sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=len(first))
        for row in data:
            sheet.append(row)
        for left, right in spans:
            sheet.merge_cells(start_row=6, start_column=left+1, end_row=6, end_column=right+1)
        for col in [0, 11, 12] + ([13] if review else []) + [len(first)-1]:
            sheet.merge_cells(start_row=6, start_column=col+1, end_row=7, end_column=col+1)
        from openpyxl.utils import get_column_letter
        for col in range(1, len(first)+1):
            sheet.column_dimensions[get_column_letter(col)].width = 14 if col == 1 else 12
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.font = Font(name='Arial', size=10, bold=cell.row <= 7 or cell.row == sheet.max_row)
                if cell.row >= 6:
                    cell.border = Border(*( [Side(style='thin', color='808080')] * 4 ))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.row_dimensions[7].height = 45
        sheet.row_dimensions[5].height = 30
        sheet.freeze_panes = 'B8'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:7'
        book.save(out)
        media = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        styles = getSampleStyleSheet()
        styles['Normal'].fontSize = 6
        styles['Normal'].leading = 8
        styles['Normal'].alignment = 1
        for i in (0, 1):
            data[i] = [Paragraph(str(value).replace('’', "'"), styles['Normal']) for value in data[i]]
        table = Table(data, colWidths=[55] + [(785-55)/(len(first)-1)]*(len(first)-1), repeatRows=2)
        commands = [('GRID', (0,0), (-1,-1), .4, colors.grey), ('ALIGN', (0,0), (-1,-1), 'CENTER'), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('FONTSIZE', (0,0), (-1,-1), 7), ('TOPPADDING', (0,0), (-1,-1), 3), ('BOTTOMPADDING', (0,0), (-1,-1), 3), ('BACKGROUND', (0,0), (-1,1), colors.HexColor('#edf2ef')), ('FONTNAME',(0,-1),(-1,-1),'Helvetica-Bold')]
        commands += [('SPAN', (left,0), (right,0)) for left,right in spans]
        commands += [('SPAN',(col,0),(col,1)) for col in [0,11,12] + ([13] if review else []) + [len(first)-1]]
        table.setStyle(TableStyle(commands))
        SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=25, rightMargin=25, topMargin=15, bottomMargin=15).build([Paragraph('DEPARTMENT OF RADIOLOGY & IMAGING<br/>CMH DHAKA<br/>' + title + '<br/>MRI CENTRE', styles['Normal']), Paragraph(note, styles['Normal']), table])
        media = 'application/pdf'
    return Response(out.getvalue(), media_type=media, headers={'Content-Disposition': f'attachment; filename="mri-summary-{period}-{basis}.{format}"'})
